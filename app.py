"""
=============================================================================
Universal Prospecting Platform (V9.0 "Turbo Core")
=============================================================================
Description: Local B2B lead generation and outreach engine (Python + Streamlit +
             Playwright + SQLite). Everything runs and is stored on your computer.

             V9.0: parallel Google Maps engine with automatic CAPTCHA/rate-limit
             handling, cascading enrichment in a thread pool (homepage, discovered
             contact pages, JavaScript websites, Facebook, search engine), email
             confidence + MX verification, international phone/WhatsApp detection,
             English/Spanish interface with help on every field, formatted Excel
             export, email preview and test sends, one-click installer.

Author:      Código Origami - Alejandro Moreno
Repository:  https://github.com/Codigo-origami
License:     MIT License
=============================================================================
"""

__author__ = "Codigo Origami - Alejandro Moreno"
__version__ = "9.0.1"

import os
import io
import sys
import time
import types
import random
import re
import datetime
import uuid
import sqlite3
import hashlib
import json
import unicodedata
import math
import threading
import asyncio
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode, urljoin, unquote, quote_plus

import requests
import pandas as pd
from bs4 import BeautifulSoup
from playwright.async_api import async_playwright
import streamlit as st
import streamlit.components.v1 as components
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

import phonenumbers
from phonenumbers import PhoneNumberType
try:
    import dns.resolver
    HAS_DNS = True
except ImportError:
    HAS_DNS = False

if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

st.set_page_config(page_title="Universal Prospector | Código Origami · Alejandro Moreno", layout="wide")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_NAME = os.path.join(BASE_DIR, "origami_leads.db")

# ============================================================================
# 1. CORE UTILITIES & NORMALIZATION (Must precede DB init)
# ============================================================================


def normalize_address(address):
    if not address or address == "Not available": return ""
    norm = unicodedata.normalize("NFKC", str(address)).casefold()
    return "".join(ch for ch in norm if ch.isalnum() or ch.isspace()).strip()

def normalize_location(loc_string):
    if not loc_string: return "unknown"
    return "|".join([unicodedata.normalize("NFKC", p.strip()).casefold() for p in loc_string.split(',') if p.strip()])

def parse_weighted_keywords(kw_text):
    kws = []
    for line in re.split(r'[\n,]+', kw_text):
        line = line.strip()
        if not line: continue
        weight = 10
        if ':' in line:
            parts = line.rsplit(':', 1)
            try: weight = max(1, min(int(parts[1].strip()), 30)); word_part = parts[0].strip().lower()
            except Exception: word_part = line.lower()
        else: word_part = line.lower()
        aliases = [alias.strip() for alias in word_part.split('|') if alias.strip()]
        if aliases: kws.append({"aliases": aliases, "weight": weight})
    return kws

def parse_flat_keywords(kw_text):
    kws = []
    for line in re.split(r'[\n,]+', kw_text):
        line = line.strip().lower()
        if not line: continue
        aliases = [alias.strip() for alias in line.split('|') if alias.strip()]
        if aliases: kws.append(aliases)
    return kws

def get_icp_hash(req_kws, strong_kws, bonus_kws, excl_kws):
    icp_state = {"req": req_kws, "strong": strong_kws, "bonus": bonus_kws, "excl": excl_kws}
    icp_str = json.dumps(icp_state, sort_keys=True)
    return hashlib.sha256(icp_str.encode()).hexdigest()[:12]

def is_social_domain(hostname, domain):
    hostname = (hostname or "").lower().rstrip(".")
    return hostname == domain or hostname.endswith("." + domain)

def email_matches_domain(email, domain):
    if not email or not domain: return False
    try:
        email_domain = email.rsplit("@", 1)[1].lower().rstrip(".")
        return email_domain == domain.lower().rstrip(".")
    except Exception: return False

def clean_url(url):
    if not url or url == "Not available": return "Not available"
    try:
        parts = urlsplit(url)
        params = dict(parse_qsl(parts.query, keep_blank_values=True))
        if 'google.' in parts.netloc or 'duckduckgo.' in parts.netloc or 'bing.' in parts.netloc:
            if "uddg" in params: return unquote(params["uddg"])
            if "q" in params: 
                candidate = unquote(params["q"])
                if candidate.startswith(("http://", "https://")): return candidate
        return url
    except Exception: return url

def normalize_company_name(name):
    if not name or name == "Unknown": return ""
    s = name.casefold().strip()
    s = "".join(ch if ch.isalnum() or ch.isspace() else " " for ch in s)
    s = re.sub(r"\b(ltd|limited|llc|inc|corp|corporation|co|company|gmbh|sa|sas|sl)\b", " ", s, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", s).strip()

def normalize_domain(url):
    if not url or url == "Not available": return ""
    if not url.startswith("http"): url = "http://" + url
    try:
        domain = urlsplit(url).hostname
        if domain:
            domain = domain.lower()
            if domain.startswith("www."): domain = domain[4:]
            SOCIAL_DOMAINS = {"facebook.com", "instagram.com", "linkedin.com", "twitter.com", "x.com"}
            if any(is_social_domain(domain, social) for social in SOCIAL_DOMAINS):
                return ""
            return domain
    except Exception: pass
    return ""

def normalize_phone(phone):
    if not phone or phone == "Not available": return "Not available"
    cleaned = re.sub(r"[^\d\+]", "", phone)
    return cleaned if cleaned else "Not available"

def extract_place_id(url):
    match = re.search(r'!1s(0x[0-9a-fA-F]+:0x[0-9a-fA-F]+)', url)
    return match.group(1) if match else "Unknown"

def add_hl_param(url, hl="en"):
    parts = urlsplit(url)
    query = dict(parse_qsl(parts.query, keep_blank_values=True))
    query["hl"] = hl
    return urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(query), parts.fragment))


COUNTRY_ISO = {
    "Afghanistan": "AF", "Albania": "AL", "Algeria": "DZ", "Andorra": "AD", "Angola": "AO",
    "Antigua and Barbuda": "AG", "Argentina": "AR", "Armenia": "AM", "Australia": "AU",
    "Austria": "AT", "Azerbaijan": "AZ", "Bahamas": "BS", "Bahrain": "BH", "Bangladesh": "BD",
    "Barbados": "BB", "Belarus": "BY", "Belgium": "BE", "Belize": "BZ", "Benin": "BJ", "Bhutan": "BT",
    "Bolivia": "BO", "Bosnia and Herzegovina": "BA", "Botswana": "BW", "Brazil": "BR", "Brunei": "BN",
    "Bulgaria": "BG", "Burkina Faso": "BF", "Burundi": "BI", "Cabo Verde": "CV", "Cambodia": "KH",
    "Cameroon": "CM", "Canada": "CA", "Central African Republic": "CF", "Chad": "TD", "Chile": "CL",
    "China": "CN", "Colombia": "CO", "Comoros": "KM", "Congo (Republic of the)": "CG",
    "Costa Rica": "CR", "Croatia": "HR", "Cuba": "CU", "Cyprus": "CY", "Czechia": "CZ",
    "Democratic Republic of the Congo": "CD", "Denmark": "DK", "Djibouti": "DJ", "Dominica": "DM",
    "Dominican Republic": "DO", "Ecuador": "EC", "Egypt": "EG", "El Salvador": "SV",
    "Equatorial Guinea": "GQ", "Eritrea": "ER", "Estonia": "EE", "Eswatini": "SZ", "Ethiopia": "ET",
    "Fiji": "FJ", "Finland": "FI", "France": "FR", "Gabon": "GA", "Gambia": "GM", "Georgia": "GE",
    "Germany": "DE", "Ghana": "GH", "Greece": "GR", "Grenada": "GD", "Guatemala": "GT", "Guinea": "GN",
    "Guinea-Bissau": "GW", "Guyana": "GY", "Haiti": "HT", "Honduras": "HN", "Hungary": "HU",
    "Iceland": "IS", "India": "IN", "Indonesia": "ID", "Iran": "IR", "Iraq": "IQ", "Ireland": "IE",
    "Israel": "IL", "Italy": "IT", "Ivory Coast": "CI", "Jamaica": "JM", "Japan": "JP", "Jordan": "JO",
    "Kazakhstan": "KZ", "Kenya": "KE", "Kiribati": "KI", "Kosovo": "XK", "Kuwait": "KW",
    "Kyrgyzstan": "KG", "Laos": "LA", "Latvia": "LV", "Lebanon": "LB", "Lesotho": "LS", "Liberia": "LR",
    "Libya": "LY", "Liechtenstein": "LI", "Lithuania": "LT", "Luxembourg": "LU", "Madagascar": "MG",
    "Malawi": "MW", "Malaysia": "MY", "Maldives": "MV", "Mali": "ML", "Malta": "MT",
    "Marshall Islands": "MH", "Mauritania": "MR", "Mauritius": "MU", "Mexico": "MX", "Micronesia": "FM",
    "Moldova": "MD", "Monaco": "MC", "Mongolia": "MN", "Montenegro": "ME", "Morocco": "MA",
    "Mozambique": "MZ", "Myanmar": "MM", "Namibia": "NA", "Nauru": "NR", "Nepal": "NP",
    "Netherlands": "NL", "New Zealand": "NZ", "Nicaragua": "NI", "Niger": "NE", "Nigeria": "NG",
    "North Korea": "KP", "North Macedonia": "MK", "Norway": "NO", "Oman": "OM", "Pakistan": "PK",
    "Palau": "PW", "Palestine State": "PS", "Panama": "PA", "Papua New Guinea": "PG", "Paraguay": "PY",
    "Peru": "PE", "Philippines": "PH", "Poland": "PL", "Portugal": "PT", "Qatar": "QA", "Romania": "RO",
    "Russia": "RU", "Rwanda": "RW", "Saint Kitts and Nevis": "KN", "Saint Lucia": "LC",
    "Saint Vincent and the Grenadines": "VC", "Samoa": "WS", "San Marino": "SM",
    "Sao Tome and Principe": "ST", "Saudi Arabia": "SA", "Senegal": "SN", "Serbia": "RS",
    "Seychelles": "SC", "Sierra Leone": "SL", "Singapore": "SG", "Slovakia": "SK", "Slovenia": "SI",
    "Solomon Islands": "SB", "Somalia": "SO", "South Africa": "ZA", "South Korea": "KR",
    "South Sudan": "SS", "Spain": "ES", "Sri Lanka": "LK", "Sudan": "SD", "Suriname": "SR",
    "Sweden": "SE", "Switzerland": "CH", "Syria": "SY", "Taiwan": "TW", "Tajikistan": "TJ",
    "Tanzania": "TZ", "Thailand": "TH", "Timor-Leste": "TL", "Togo": "TG", "Tonga": "TO",
    "Trinidad and Tobago": "TT", "Tunisia": "TN", "Turkey": "TR", "Turkmenistan": "TM", "Tuvalu": "TV",
    "Uganda": "UG", "Ukraine": "UA", "United Arab Emirates": "AE", "United Kingdom": "GB",
    "United States of America": "US", "Uruguay": "UY", "Uzbekistan": "UZ", "Vanuatu": "VU",
    "Vatican City": "VA", "Venezuela": "VE", "Vietnam": "VN", "Yemen": "YE", "Zambia": "ZM",
    "Zimbabwe": "ZW"
}
COUNTRY_ISO_LOWER = {k.lower(): v for k, v in COUNTRY_ISO.items()}


def get_high_signal_text(soup):
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "svg", "form"]):
        tag.decompose()
    garbage_pattern = re.compile(r"cookie|privacy|terms|login|menu", re.I)
    for tag in soup.find_all(attrs={"class": garbage_pattern}): tag.decompose()
    for tag in soup.find_all(attrs={"id": garbage_pattern}): tag.decompose()

    parts = []
    if soup.title and soup.title.string: parts.append(soup.title.string)
    meta_desc = soup.find('meta', attrs={'name': 'description'})
    if meta_desc and meta_desc.get('content'): parts.append(meta_desc['content'])
    for tag in ['h1', 'h2', 'h3']:
        for el in soup.find_all(tag):
            parts.append(el.get_text(separator=' ', strip=True))
            
    body_text = soup.get_text(separator=' ', strip=True)
    signal_body = body_text[:1000] + " " + body_text[-500:] if len(body_text) > 1500 else body_text
    parts.append(signal_body)
    return " ".join(parts).lower()



# ============================================================================
# 3. EMAIL / PHONE QUALITY (V9)
# ============================================================================

NA = "Not available"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/126.0.0.0 Safari/537.36")
EMAIL_RE = re.compile(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9-]+(?:\.[a-zA-Z0-9-]+)*\.[a-zA-Z]{2,24}')
FREE_PROVIDERS = {
    "gmail.com", "googlemail.com", "hotmail.com", "outlook.com", "live.com", "msn.com", "yahoo.com",
    "ymail.com", "rocketmail.com", "icloud.com", "me.com", "mac.com", "aol.com", "protonmail.com",
    "proton.me", "gmx.com", "gmx.net", "mail.com", "zoho.com", "yandex.com"
}
BAD_EMAIL_PARTS = (
    "sentry", "wixpress", "example.", "yourdomain", "your-domain", "domain.com", "email.com", "mysite",
    "yoursite", "yourcompany", "company.com", "test.com", "schema.org", "w3.org", "godaddy", "polyfill",
    "github", "cloudflare", "gravatar", "wordpress", "noreply", "no-reply", "donotreply", "do-not-reply",
    "jquery", "bootstrap", "@2x", "@3x", "ingest.", "squarespace.com", "shopify.com", "wix.com",
    "website.com", "name@", "user@", "username@", "johndoe", "firstname", "@sample", "placeholder",
    "u00", "%20", "duckduckgo", "error-lite", "@bing.", "@google.", "@sentry"
)
ROLE_SCORES = [("purchasing@", 40), ("procurement@", 40), ("compras@", 40), ("import@", 35),
               ("sales@", 25), ("ventas@", 25), ("comercial@", 25), ("marketing@", 15),
               ("info@", 15), ("contact@", 15), ("contacto@", 15), ("office@", 12), ("hello@", 12),
               ("admin@", 5)]
CONF_RANK = {"HIGH": 3, "MEDIUM": 2, "LOW": 1, "NONE": 0}
GENERIC_NAME_WORDS = {
    "the", "and", "company", "group", "international", "global", "services", "service", "solutions",
    "trading", "store", "shop", "center", "centre", "sales", "limited", "enterprise", "enterprises",
    "industries", "industry", "import", "imports", "export", "auto", "autos", "parts", "motors", "motor"
}


def clean_email(e):
    e = unquote(str(e or "")).strip().strip(".,;:'\"<>()[]").lower()
    return e.replace("mailto:", "")


def is_valid_email(email):
    email = clean_email(email)
    if not email or email.count("@") != 1: return False
    local, domain = email.split("@")
    if not (1 <= len(local) <= 64) or len(email) > 254: return False
    if ".." in email or local.startswith(".") or local.endswith("."): return False
    if not re.fullmatch(r'[a-z0-9.-]+\.[a-z]{2,24}', domain): return False
    if re.search(r'\.(png|jpe?g|gif|webp|svg|css|js|ico|bmp|tiff?|mp4|pdf)$', email): return False
    if re.fullmatch(r'[a-f0-9]{20,}', local): return False          # tracking hashes
    if re.search(r'@[0-9-]+\.[a-z0-9.-]+$', email): return False
    return not any(b in email for b in BAD_EMAIL_PARTS)


def decode_cfemail(encoded):
    try:
        r = int(encoded[:2], 16)
        return ''.join(chr(int(encoded[i:i + 2], 16) ^ r) for i in range(2, len(encoded), 2))
    except Exception:
        return ""


def role_score(email):
    score = next((s for p, s in ROLE_SCORES if email.startswith(p)), 0)
    if "privacy" in email or "gdpr" in email: score -= 30
    if any(w in email for w in ("support", "careers", "jobs", "hr@", "rrhh")): score -= 15
    return score


def get_best_email(emails):
    emails = [clean_email(e) for e in emails if is_valid_email(e)]
    if not emails: return NA
    return max(emails, key=role_score)


def is_free_mail(domain):
    return domain in FREE_PROVIDERS or bool(re.match(r'^(hotmail|yahoo|outlook|live)\.[a-z.]+$', domain))


def company_tokens(name):
    return {t for t in normalize_company_name(name).split() if len(t) >= 4 and t not in GENERIC_NAME_WORDS}


def classify_email(email, source, site_domain, tokens):
    """How sure we are that the email belongs to THIS company."""
    domain = email.split("@")[-1]
    token_hit = any(t in email for t in tokens)
    if site_domain and (domain == site_domain or domain.endswith("." + site_domain)
                        or site_domain.endswith("." + domain)):
        return "HIGH"
    if source == "Search Engine":
        return "MEDIUM" if token_hit else "LOW"
    if is_free_mail(domain):
        return "MEDIUM"           # small businesses often use gmail/hotmail on their own site
    if token_hit:
        return "HIGH"
    return "LOW"                  # another company's domain (web agency, supplier…)


_mx_cache = {}


def domain_accepts_mail(domain):
    if domain in _mx_cache: return _mx_cache[domain]
    ok = True
    if HAS_DNS and not is_free_mail(domain):
        try:
            dns.resolver.resolve(domain, "MX", lifetime=6)
        except dns.resolver.NXDOMAIN:
            ok = False
        except dns.resolver.NoAnswer:
            try:
                dns.resolver.resolve(domain, "A", lifetime=6)
            except Exception:
                ok = False
        except Exception:
            ok = True     # DNS timeout: don't punish the lead
    _mx_cache[domain] = ok
    return ok


def choose_best_email(candidates, site_domain, tokens, verify_mx):
    """candidates: list of (email, source, source_url). Returns (email, source, url, confidence)."""
    seen, scored = set(), []
    for email, source, url in candidates:
        email = clean_email(email)
        if email in seen or not is_valid_email(email): continue
        seen.add(email)
        conf = classify_email(email, source, site_domain, tokens)
        if source == "Search Engine" and conf == "LOW": continue      # random result: discard
        scored.append((CONF_RANK[conf], role_score(email), email, source, url, conf))
    scored.sort(reverse=True)
    for _, _, email, source, url, conf in scored[:6]:
        if not verify_mx or domain_accepts_mail(email.split("@")[-1]):
            return email, source, url, conf
    return NA, NA, NA, "NONE"


def best_conf(candidates, site_domain, tokens):
    return max((CONF_RANK[classify_email(clean_email(e), s, site_domain, tokens)]
                for e, s, _ in candidates if is_valid_email(e)), default=0)


# ── Countries and phones ──────────────────────────────────────────────────
COUNTRY_ALIASES = {
    "españa": "ES", "méxico": "MX", "mexico": "MX", "perú": "PE", "peru": "PE", "panamá": "PA",
    "república dominicana": "DO", "republica dominicana": "DO", "estados unidos": "US", "usa": "US",
    "eeuu": "US", "united states": "US", "uk": "GB", "reino unido": "GB", "brasil": "BR",
    "alemania": "DE", "francia": "FR", "italia": "IT", "japón": "JP", "japon": "JP", "china": "CN",
    "marruecos": "MA", "sudáfrica": "ZA", "sudafrica": "ZA", "nueva zelanda": "NZ", "canadá": "CA",
    "chipre": "CY", "grecia": "GR", "turquía": "TR", "turquia": "TR", "egipto": "EG", "kenia": "KE",
    "trinidad y tobago": "TT", "jamaica": "JM", "costa de marfil": "CI", "côte d'ivoire": "CI",
    "emiratos árabes unidos": "AE", "uae": "AE", "arabia saudí": "SA", "arabia saudita": "SA",
    "países bajos": "NL", "holanda": "NL", "bélgica": "BE", "suiza": "CH", "suecia": "SE",
    "noruega": "NO", "polonia": "PL", "rusia": "RU", "corea del sur": "KR", "filipinas": "PH",
    "tailandia": "TH", "vietnam": "VN", "malasia": "MY", "singapur": "SG", "nigeria": "NG",
    "ghana": "GH", "tanzania": "TZ", "uganda": "UG", "zambia": "ZM", "zimbabue": "ZW",
    "mozambique": "MZ", "angola": "AO", "namibia": "NA", "botsuana": "BW", "camerún": "CM",
    "argelia": "DZ", "túnez": "TN", "líbano": "LB", "jordania": "JO", "irak": "IQ", "irán": "IR",
    "pakistán": "PK", "india": "IN", "bangladés": "BD", "sri lanka": "LK", "nepal": "NP",
    "mongolia": "MN", "kazajistán": "KZ", "uzbekistán": "UZ", "georgia": "GE", "armenia": "AM",
    "azerbaiyán": "AZ", "ucrania": "UA", "rumanía": "RO", "bulgaria": "BG", "serbia": "RS",
    "croacia": "HR", "albania": "AL", "portugal": "PT", "irlanda": "IE", "australia": "AU",
    "fiyi": "FJ", "papúa nueva guinea": "PG", "guyana": "GY", "surinam": "SR", "belice": "BZ",
    "barbados": "BB", "bahamas": "BS", "haití": "HT", "haiti": "HT", "cuba": "CU",
}


def region_for_location(location):
    """'Santiago, Chile' → 'CL' (used to read local phone numbers correctly)."""
    parts = [p.strip() for p in str(location or "").split(",") if p.strip()]
    if not parts: return None
    country = parts[-1].lower()
    if country in COUNTRY_ALIASES: return COUNTRY_ALIASES[country]
    for name, iso in COUNTRY_ISO_LOWER.items():
        if country == name: return iso
    return None


def parse_phone(raw, region):
    if not raw or raw == NA: return None
    raw = str(raw).strip()
    digits = re.sub(r'\D', '', raw)
    if len(digits) < 7: return None
    attempts = [raw, digits] if raw.startswith("+") else [raw, "+" + digits]
    for a in attempts:
        try:
            n = phonenumbers.parse(a, region)
            if phonenumbers.is_valid_number(n): return n
        except phonenumbers.NumberParseException:
            continue
    return None


def format_phone(raw, region):
    n = parse_phone(raw, region)
    if n: return phonenumbers.format_number(n, phonenumbers.PhoneNumberFormat.INTERNATIONAL)
    cleaned = re.sub(r'[^\d\+\-\s\(\)]', '', str(raw or "")).strip()
    return cleaned or NA


def to_e164(raw, region):
    n = parse_phone(raw, region)
    return phonenumbers.format_number(n, phonenumbers.PhoneNumberFormat.E164) if n else None


def is_mobile(raw, region):
    n = parse_phone(raw, region)
    return bool(n) and phonenumbers.number_type(n) == PhoneNumberType.MOBILE


# ============================================================================
# 4. WEBSITE ENRICHMENT (plain HTTP in a thread pool, browser only when needed)
# ============================================================================
SOCIAL_HOSTS = ("facebook.com", "fb.com", "instagram.com", "linkedin.com", "twitter.com", "x.com",
                "tiktok.com", "youtube.com")
SOCIAL_JUNK = ("sharer", "/share", "/dialog/", "/plugins/", "/policies", "/help", "/login", "/legal",
               "/privacy", "/tos", "/intent/", "/hashtag/", "/watch", "/tr?", "/ads", "/oauth")
CONTACT_HINTS = [("contact", 10), ("contacto", 10), ("contactenos", 10), ("contactanos", 10),
                 ("contato", 10), ("kontakt", 10), ("contatti", 10), ("get-in-touch", 9), ("enquir", 8),
                 ("inquir", 8), ("reach-us", 8), ("find-us", 7), ("location", 5), ("about", 6),
                 ("nosotros", 6), ("quienes", 6), ("sobre", 5), ("acerca", 5), ("empresa", 4),
                 ("impressum", 7), ("aviso-legal", 6), ("legal", 3), ("team", 3)]
WA_TEXT_RE = re.compile(r'(\bwhats\s?app\b|\bwa\b|\bmobile\b|\bcell\b|\bcelular\b|\bm[óo]vil\b|📱)[\s:.\-]*'
                        r'(\+?[\d\s\-\(\)]{8,20})', re.IGNORECASE)


def host_of(url):
    if not url or url == NA: return ""
    if not url.startswith("http"): url = "http://" + url
    try:
        h = (urlsplit(url).hostname or "").lower()
        return h[4:] if h.startswith("www.") else h
    except Exception:
        return ""


def is_social(url):
    h = host_of(url)
    return any(h == d or h.endswith("." + d) for d in SOCIAL_HOSTS)


def clean_social(url):
    url = clean_url(url)
    if not url or url == NA or not is_social(url): return ""
    if any(j in url.lower() for j in SOCIAL_JUNK): return ""
    if "profile.php" in url.lower(): return url.split("&")[0]
    return url.split("?")[0].rstrip("/")


def split_socials(urls):
    out = {"facebook": NA, "instagram": NA, "linkedin": NA}
    for u in urls:
        c = clean_social(u)
        if not c or not urlsplit(c).path.strip("/"): continue
        h = host_of(c)
        key = "facebook" if ("facebook" in h or h == "fb.com") else "instagram" if "instagram" in h \
            else "linkedin" if "linkedin" in h else None
        if key and out[key] == NA: out[key] = c
    return out


def wa_from_text(text):
    return [(m.group(2), "WA" if re.match(r'whats|wa', m.group(1), re.I) else None)
            for m in WA_TEXT_RE.finditer(text or "")]


def emails_in_text(text):
    out = set()
    for m in EMAIL_RE.findall(text or ""):
        m = re.sub(r'(\.[a-z]{2,})[A-Z][A-Za-z]*$', r'\1', m)      # "info@x.comPhone" → "info@x.com"
        e = clean_email(m)
        if is_valid_email(e): out.add(e)
    return out


def deobfuscate(text):
    for a, b in (("[at]", "@"), ("(at)", "@"), ("{at}", "@"), (" [at] ", "@"), ("&#64;", "@"),
                 ("&#x40;", "@"), ("%40", "@"), ("[dot]", "."), ("(dot)", "."), ("{dot}", ".")):
        text = text.replace(a, b)
    return text


def extract_from_html(html, base_url=""):
    html = deobfuscate(html or "")
    soup = BeautifulSoup(html, "html.parser")
    signal = get_high_signal_text(BeautifulSoup(html, "html.parser"))
    for tag in soup(["script", "style", "noscript"]):
        if tag.name == "script" and tag.get("type") == "application/ld+json": continue
        tag.decompose()
    text = soup.get_text(" ", strip=True)

    emails = set()
    for tag in soup.select("[data-cfemail]"):
        e = clean_email(decode_cfemail(tag.get("data-cfemail", "")))
        if is_valid_email(e): emails.add(e)
    for a in soup.select('a[href^="mailto:"], a[href^="MAILTO:"]'):
        e = clean_email(a["href"].split(":", 1)[1].split("?")[0])
        if is_valid_email(e): emails.add(e)
    emails |= emails_in_text(html) | emails_in_text(text)

    wa, tels, socials, links = [], [], set(), []
    base_host = host_of(base_url)
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        low = href.lower()
        if "wa.me/" in low or "whatsapp.com/send" in low or "whatsapp://send" in low:
            m = re.search(r'(?:wa\.me/|phone=)(\+?\d{7,15})', low)
            if m: wa.append((m.group(1), href))
            continue
        if low.startswith("tel:"):
            tels.append(unquote(href[4:])); continue
        absolute = urljoin(base_url, href) if base_url else href
        if is_social(absolute): socials.add(absolute)
        elif base_host and host_of(absolute) == base_host:
            links.append((absolute, a.get_text(" ", strip=True).lower()))
    return {"emails": emails, "wa": wa + wa_from_text(text), "tels": tels, "socials": socials,
            "links": links, "text_len": len(text), "signal": signal}


def contact_links(links, limit=4):
    scored = {}
    for url, label in links:
        path = urlsplit(url).path.lower()
        if re.search(r'\.(pdf|jpe?g|png|zip|docx?)$', path): continue
        score = max((s for k, s in CONTACT_HINTS if k in path + " " + label), default=0)
        if score:
            key = url.split("#")[0].rstrip("/")
            scored[key] = max(score, scored.get(key, 0))
    return [u for u, _ in sorted(scored.items(), key=lambda x: -x[1])][:limit]


def looks_js_rendered(html, info):
    markers = ("__NEXT_DATA__", 'id="root"', 'id="app"', "ng-version", "data-reactroot",
               "window.__NUXT__", "static.wixstatic.com", "elementor")
    return info["text_len"] < 600 or any(m in html for m in markers)


def new_session():
    s = requests.Session()
    s.headers.update({"User-Agent": UA, "Accept-Language": "en-US,en;q=0.9,es;q=0.8",
                      "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"})
    return s


def fetch_html(session, url, timeout=9):
    """(html, final_url) — html is None when missing or blocked (then a real browser may help)."""
    if not url.startswith("http"): url = "https://" + url
    for attempt in dict.fromkeys((url, url.replace("https://", "http://", 1))):
        try:
            r = session.get(attempt, timeout=timeout, allow_redirects=True)
        except requests.RequestException:
            continue
        if r.ok and "html" in r.headers.get("Content-Type", "").lower():
            return r.text[:2_500_000], r.url
        return None, r.url
    return None, url


def site_static(website, deep):
    """Homepage (+ discovered contact pages when deep). Runs in a worker thread."""
    acc = {"emails": [], "wa": [], "tels": [], "socials": set(), "signal": "", "contact_urls": [],
           "needs_render": False, "final_domain": normalize_domain(website), "home_url": website, "errors": []}
    session = new_session()
    html, final_url = fetch_html(session, website)
    if html is None:
        acc["needs_render"] = True
        acc["errors"].append("Homepage blocked/unreachable")
        return acc
    acc["home_url"] = final_url
    acc["final_domain"] = normalize_domain(final_url) or acc["final_domain"]
    info = extract_from_html(html, final_url)
    acc["emails"] += [(e, "Website", final_url) for e in info["emails"]]
    acc["wa"] += info["wa"]; acc["tels"] += info["tels"]; acc["socials"] |= info["socials"]
    acc["signal"] = info["signal"]
    site = acc["final_domain"]
    candidates = contact_links(info["links"]) or \
        [urljoin(final_url, p) for p in ("/contact", "/contact-us", "/contacto", "/about")]
    acc["contact_urls"] = candidates
    if deep:
        for url in candidates:
            if site and any(e.split("@")[-1].endswith(site) for e, _, _ in acc["emails"]): break
            sub, sub_url = fetch_html(session, url, timeout=7)
            if not sub: continue
            i2 = extract_from_html(sub, sub_url)
            acc["emails"] += [(e, "Website (Contact)", sub_url) for e in i2["emails"]]
            acc["wa"] += i2["wa"]; acc["tels"] += i2["tels"]; acc["socials"] |= i2["socials"]
            acc["signal"] += " " + i2["signal"][:1500]
    if not acc["emails"] and looks_js_rendered(html, info):
        acc["needs_render"] = True
    return acc


# ── DuckDuckGo: one request at a time; stops for the run if it starts blocking ──
_ddg_lock = threading.Lock()
_ddg_state = {"last": 0.0, "blocked": False}


def ddg_search(query):
    if _ddg_state["blocked"]: return ""
    if not _ddg_lock.acquire(timeout=60): return ""
    try:
        wait = 3.0 - (time.time() - _ddg_state["last"])
        if wait > 0: time.sleep(wait + random.uniform(0, 1.5))
        _ddg_state["last"] = time.time()
        res = requests.post("https://lite.duckduckgo.com/lite/", data={"q": query},
                            headers={"User-Agent": UA}, timeout=12)
    except requests.RequestException:
        return ""
    finally:
        _ddg_lock.release()
    text = res.text or ""
    if res.status_code != 200 or "error-lite" in text or "anomaly" in text.lower():
        _ddg_state["blocked"] = True
        return ""
    return text


def search_emails(company, location):
    html = ddg_search(f'"{company}" {location} email')
    if not html: return []
    return sorted(emails_in_text(BeautifulSoup(html, "html.parser").get_text(" ")) | emails_in_text(html))


def search_socials(company, location):
    html = ddg_search(f'"{company}" {location} facebook OR instagram OR linkedin')
    if not html: return {}
    soup = BeautifulSoup(html, "html.parser")
    return split_socials(clean_url(a["href"]) for a in soup.find_all("a", href=True))



# --- STRICT IDENTITY ALIASING ---

def generate_identity_aliases(maps_id, domain, phone, norm_name, norm_address):
    aliases = []
    if maps_id and maps_id != "Unknown": aliases.append(f"maps:{maps_id}")
    if domain and norm_address: aliases.append(f"dom_addr:{domain}|{norm_address}")
    if phone and phone != "Not available" and norm_address: aliases.append(f"ph_addr:{phone}|{norm_address}")
    if norm_name and norm_address: aliases.append(f"name_addr:{norm_name}|{norm_address}")
    return list(set(aliases))

def resolve_lead_identity(db_conn, aliases):
    if not aliases: return None, "NO_IDENTITY", []
    
    cur = db_conn.cursor()
    placeholders = ', '.join(['?'] * len(aliases))
    cur.execute(f"SELECT DISTINCT lead_id FROM lead_identity_aliases WHERE alias_key IN ({placeholders})", tuple(aliases))
    results = [row[0] for row in cur.fetchall()]
    
    if len(results) == 1:
        return results[0], "EXISTING", []
    elif len(results) > 1:
        return None, "IDENTITY_CONFLICT", results
    
    return f"ORI-{str(uuid.uuid4())[:8].upper()}", "NEW", []

def log_identity_conflict(db_conn, run_id, aliases, conflict_ids, company_name, maps_id, domain, phone, address):
    conflict_id = f"CONF-{uuid.uuid4().hex[:8].upper()}"
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    id_1 = conflict_ids[0] if len(conflict_ids) > 0 else "UNKNOWN"
    id_2 = conflict_ids[1] if len(conflict_ids) > 1 else "UNKNOWN"
    
    db_conn.execute("""
        INSERT INTO identity_conflicts 
        (conflict_id, run_id, alias_key, lead_id_1, lead_id_2, company_name, maps_id, domain, phone, address, detected_at, resolution_status) 
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (conflict_id, run_id, json.dumps(aliases), id_1, id_2, company_name, maps_id, domain, phone, address, now, "PENDING"))

# ============================================================================
# 2. DATABASE LAYER
# ============================================================================

def get_db():
    conn = sqlite3.connect(DB_NAME, timeout=15)
    conn.execute("PRAGMA busy_timeout=10000")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn

def add_column_if_missing(conn, table, column, column_type):
    existing = {row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}
    if column not in existing:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {column_type}")

def ensure_db_schema():
    with get_db() as conn:
        conn.execute("PRAGMA journal_mode=WAL") 
        
        # Base Tables Creation
        conn.execute('''CREATE TABLE IF NOT EXISTS leads (lead_id TEXT PRIMARY KEY)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS lead_identity_aliases (alias_key TEXT PRIMARY KEY, lead_id TEXT NOT NULL, FOREIGN KEY(lead_id) REFERENCES leads(lead_id) ON DELETE CASCADE)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS identity_conflicts (conflict_id TEXT PRIMARY KEY)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS global_recipient_locks (recipient TEXT PRIMARY KEY)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS recipient_history (recipient TEXT PRIMARY KEY)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS lead_discoveries (run_id TEXT, lead_id TEXT, query TEXT, location TEXT, PRIMARY KEY (run_id, lead_id, query, location), FOREIGN KEY(lead_id) REFERENCES leads(lead_id) ON DELETE CASCADE)''')
        conn.execute('''CREATE TABLE IF NOT EXISTS activities (activity_id TEXT PRIMARY KEY, lead_id TEXT, FOREIGN KEY(lead_id) REFERENCES leads(lead_id) ON DELETE CASCADE)''')
        
        # Universal Schema Migration
        lead_cols = [
            ("created_at", "TEXT"), ("last_seen_at", "TEXT"), ("enriched_at", "TEXT"), ("last_attempted_at", "TEXT"), 
            ("last_contacted_at", "TEXT"), ("status_updated_at", "TEXT"), ("lock_token", "TEXT"),
            ("maps_id", "TEXT"), ("maps_url", "TEXT"), ("domain", "TEXT"), ("norm_name", "TEXT"), 
            ("norm_address", "TEXT"), ("location", "TEXT"), ("query", "TEXT"), ("company", "TEXT"), 
            ("category", "TEXT"), ("address", "TEXT"), ("website", "TEXT"), ("phone", "TEXT"), 
            ("whatsapp", "TEXT"), ("email", "TEXT"), ("email_source", "TEXT"), ("email_source_url", "TEXT"),
            ("wa_link", "TEXT"), ("facebook", "TEXT"), ("instagram", "TEXT"), ("linkedin", "TEXT"), 
            ("contactability_score", "INTEGER"), ("icp_fit_score", "INTEGER"), ("priority_score", "INTEGER"), 
            ("matched_keywords", "TEXT"), ("qual_status", "TEXT"), ("priority_badge", "TEXT"), 
            ("status", "TEXT"), ("enrichment_level", "TEXT"), ("enrich_time", "TEXT"), ("icp_hash", "TEXT"),
            ("sources_searched", "TEXT"), ("sources_found", "TEXT"), ("debug_logs", "TEXT"),
            ("email_confidence", "TEXT"), ("country", "TEXT")
        ]
        for c, t in lead_cols: add_column_if_missing(conn, "leads", c, t)
            
        act_cols = [
            ("campaign_id", "TEXT"), ("lock_token", "TEXT"), ("type", "TEXT"), ("status", "TEXT"), 
            ("recipient", "TEXT"), ("sender", "TEXT"), ("subject", "TEXT"), ("attempted_at", "TEXT"), 
            ("sent_at", "TEXT"), ("error", "TEXT")
        ]
        for c, t in act_cols: add_column_if_missing(conn, "activities", c, t)
            
        conf_cols = [
            ("run_id", "TEXT"), ("alias_key", "TEXT"), ("lead_id_1", "TEXT"), ("lead_id_2", "TEXT"), 
            ("company_name", "TEXT"), ("maps_id", "TEXT"), ("domain", "TEXT"), ("phone", "TEXT"), 
            ("address", "TEXT"), ("detected_at", "TEXT"), ("resolution_status", "TEXT")
        ]
        for c, t in conf_cols: add_column_if_missing(conn, "identity_conflicts", c, t)
            
        lock_cols = [("locked_at", "TEXT"), ("campaign_id", "TEXT"), ("lock_token", "TEXT")]
        for c, t in lock_cols: add_column_if_missing(conn, "global_recipient_locks", c, t)
            
        hist_cols = [("first_contacted_at", "TEXT"), ("last_contacted_at", "TEXT"), ("last_campaign_id", "TEXT"), ("total_attempts", "INTEGER DEFAULT 0")]
        for c, t in hist_cols: add_column_if_missing(conn, "recipient_history", c, t)
            
        disc_cols = [("discovered_at", "TEXT")]
        for c, t in disc_cols: add_column_if_missing(conn, "lead_discoveries", c, t)

        # Indexes
        conn.execute('CREATE INDEX IF NOT EXISTS idx_maps_id ON leads(maps_id)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_priority_score ON leads(priority_score DESC)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_alias_lead_id ON lead_identity_aliases(lead_id)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_leads_status ON leads(status)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_leads_qual ON leads(qual_status, priority_badge)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_activities_lead ON activities(lead_id)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_activities_camp ON activities(campaign_id)')
        conn.execute('CREATE INDEX IF NOT EXISTS idx_activities_status ON activities(status)')

def init_db():
    ensure_db_schema()

def clear_db():
    """Atomic drop with concurrent active-sends check inside a BEGIN IMMEDIATE lock."""
    with get_db() as conn:
        try:
            conn.execute("BEGIN IMMEDIATE")
            cur = conn.cursor()
            
            cur.execute("SELECT COUNT(*) FROM global_recipient_locks")
            if cur.fetchone()[0] > 0: 
                conn.rollback()
                return False
                
            cur.execute("SELECT COUNT(*) FROM leads WHERE status='Sending...'")
            if cur.fetchone()[0] > 0: 
                conn.rollback()
                return False
                
            tables = [
                "activities", "lead_discoveries", "lead_identity_aliases", 
                "global_recipient_locks", "identity_conflicts", "recipient_history", "leads"
            ]
            for table in tables:
                conn.execute(f"DROP TABLE IF EXISTS {table}")
            conn.commit()
        except Exception:
            conn.rollback()
            raise
            
    ensure_db_schema()
    return True

init_db()

def load_leads():
    with get_db() as conn:
        return pd.read_sql("SELECT * FROM leads ORDER BY priority_score DESC", conn)

def get_campaign_candidates(target_qualities, limit, skipped_ids=None):
    if not target_qualities: return pd.DataFrame()
    skipped_ids = skipped_ids or []
    
    with get_db() as conn:
        placeholders = ','.join(['?']*len(target_qualities))
        query = f"""
            SELECT lead_id, email, company, qual_status, priority_badge, status 
            FROM leads 
            WHERE qual_status='QUALIFIED' 
            AND priority_badge IN ({placeholders}) 
            AND email IS NOT NULL AND email != 'Not available' 
            AND status NOT IN ('Contacted', 'Do Not Contact', 'Replied', 'Not Interested', 'Recipient rejected', 'Sender rejected', 'Excluded', 'Not ICP', 'Sending...', 'Send outcome unknown', 'Invalid Email')
        """
        params = list(target_qualities)
        if skipped_ids:
            skip_placeholders = ','.join(['?']*len(skipped_ids))
            query += f" AND lead_id NOT IN ({skip_placeholders})"
            params.extend(skipped_ids)
            
        query += " ORDER BY priority_score DESC LIMIT ?"
        params.append(limit)
        
        return pd.read_sql(query, conn, params=tuple(params))

def upsert_lead_transaction(db_conn, lead_dict, aliases, discovery_dict):
    try:
        db_conn.execute("BEGIN IMMEDIATE")
        lead_id, id_status, conflict_ids = resolve_lead_identity(db_conn, aliases)
        
        if not lead_id: 
            if id_status == "IDENTITY_CONFLICT":
                log_identity_conflict(db_conn, discovery_dict['run_id'], aliases, conflict_ids, lead_dict['company'], lead_dict.get('maps_id'), lead_dict.get('domain'), lead_dict.get('phone'), lead_dict.get('address'))
            db_conn.commit()
            return f"Skipped: {id_status}"
            
        lead_dict['lead_id'] = lead_id
        cur = db_conn.cursor()
        
        cols = list(lead_dict.keys())
        placeholders = ', '.join(['?'] * len(cols))
        col_names = ', '.join(cols)
        
        update_clauses = []
        for col in cols:
            # created_at is ignored on update, preserving the original discovery time
            if col in ['lead_id', 'created_at']: continue
            elif col == 'status':
                # Protect hard terminal states & active locks
                update_clauses.append(f"{col} = CASE WHEN leads.status IN ('Contacted', 'Replied', 'Not Interested', 'Do Not Contact', 'Excluded', 'Sending...', 'Send outcome unknown') THEN leads.status ELSE excluded.{col} END")
            elif col in ['enriched_at', 'last_seen_at', 'last_contacted_at', 'last_attempted_at', 'status_updated_at']:
                 update_clauses.append(f"{col} = COALESCE(excluded.{col}, leads.{col})")
            else:
                update_clauses.append(f"{col} = excluded.{col}")
                
        update_clause_str = ',\n        '.join(update_clauses)
        
        query = f"""
        INSERT INTO leads ({col_names})
        VALUES ({placeholders})
        ON CONFLICT(lead_id) DO UPDATE SET
        {update_clause_str}
        """
        cur.execute(query, tuple(lead_dict[col] for col in cols))
        
        # Insert aliases AFTER the lead is safely upserted to respect FK
        for alias in aliases:
            cur.execute("SELECT lead_id FROM lead_identity_aliases WHERE alias_key=?", (alias,))
            existing = cur.fetchone()
            if existing and existing[0] != lead_id:
                raise ValueError(f"Identity collision: Alias {alias} belongs to {existing[0]}, attempted on {lead_id}")
            cur.execute("INSERT OR IGNORE INTO lead_identity_aliases (alias_key, lead_id) VALUES (?, ?)", (alias, lead_id))
            
        discovery_dict['lead_id'] = lead_id
        cur.execute('''INSERT OR IGNORE INTO lead_discoveries 
                       (run_id, lead_id, query, location, discovered_at) 
                       VALUES (?, ?, ?, ?, ?)''', 
                    (discovery_dict['run_id'], discovery_dict['lead_id'], discovery_dict['query'], 
                     discovery_dict['location'], discovery_dict['discovered_at']))
                     
        db_conn.commit()
        return "Processed"
    except Exception as e:
        db_conn.rollback()
        raise e

def lock_lead_and_create_activity(lead_id, campaign_id, recipient, sender, subject):
    with get_db() as conn:
        try:
            conn.execute("BEGIN IMMEDIATE")
            cur = conn.cursor()
            
            cur.execute("SELECT 1 FROM global_recipient_locks WHERE recipient=?", (recipient,))
            if cur.fetchone(): 
                conn.commit()
                return None 
                
            cur.execute("SELECT first_contacted_at FROM recipient_history WHERE recipient=?", (recipient,))
            row = cur.fetchone()
            if row and row[0]:
                conn.commit()
                return None
                
            lock_token = uuid.uuid4().hex
            now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            
            # Record last_attempted_at at the start of the lock to accurately track attempt timestamps
            cur.execute("""
                UPDATE leads 
                SET status='Sending...', status_updated_at=?, last_attempted_at=?, lock_token=? 
                WHERE lead_id=? AND status NOT IN ('Contacted', 'Do Not Contact', 'Replied', 'Not Interested', 'Recipient rejected', 'Sender rejected', 'Excluded', 'Not ICP', 'Sending...', 'Send outcome unknown', 'Invalid Email')
            """, (now, now, lock_token, lead_id))
            
            if cur.rowcount == 0: 
                conn.commit()
                return None
                
            cur.execute("INSERT INTO global_recipient_locks (recipient, locked_at, campaign_id, lock_token) VALUES (?, ?, ?, ?)", (recipient, now, campaign_id, lock_token))
            
            # total_attempts = number of times the system initiated contact (concurrency & preflight intents)
            cur.execute("INSERT OR IGNORE INTO recipient_history (recipient, total_attempts) VALUES (?, 0)", (recipient,))
            cur.execute("UPDATE recipient_history SET total_attempts = total_attempts + 1, last_campaign_id=? WHERE recipient=?", (campaign_id, recipient))
            
            activity_id = f"ACT-{uuid.uuid4().hex[:8].upper()}"
            cur.execute('''INSERT INTO activities 
                (activity_id, lead_id, campaign_id, lock_token, type, status, recipient, sender, subject, attempted_at, error) 
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                (activity_id, lead_id, campaign_id, lock_token, "Email", "Attempting", recipient, sender, subject, now, ""))
            
            cur.execute("SELECT email, qual_status, priority_badge, status, company FROM leads WHERE lead_id=? AND lock_token=?", (lead_id, lock_token))
            row = cur.fetchone()
            conn.commit()
            
            if not row: return None
            return {
                "email": row[0], "qual_status": row[1], "priority_badge": row[2], 
                "status": row[3], "company": row[4], "lock_token": lock_token,
                "activity_id": activity_id
            }
        except Exception:
            conn.rollback()
            raise

def update_activity_and_lead(activity_id, lead_id, act_status, lead_status, lock_token, error_msg="", recipient=None):
    with get_db() as conn:
        try:
            conn.execute("BEGIN IMMEDIATE")
            cur = conn.cursor()
            now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            
            time_updates = "status_updated_at=?, lock_token=NULL"
            lead_params = [lead_status, now]
            if act_status == "Contacted":
                time_updates += ", last_contacted_at=?"
                lead_params.append(now)
                
            lead_params.extend([lead_id, lock_token])
            cur.execute(f"UPDATE leads SET status=?, {time_updates} WHERE lead_id=? AND lock_token=? AND status='Sending...'", tuple(lead_params))
            
            if cur.rowcount != 1:
                raise RuntimeError(f"Lost send lock for lead {lead_id} (Token: {lock_token}). Update aborted to protect DB integrity.")
            
            sent_update = ", sent_at=?" if act_status == "Contacted" else ""
            act_params = [act_status, error_msg]
            if act_status == "Contacted": act_params.append(now)
            act_params.append(activity_id)
            cur.execute(f"UPDATE activities SET status=?, error=?{sent_update} WHERE activity_id=?", tuple(act_params))
            
            if recipient:
                if act_status == "Contacted":
                    cur.execute("UPDATE recipient_history SET first_contacted_at = COALESCE(first_contacted_at, ?), last_contacted_at = ? WHERE recipient = ?", (now, now, recipient))
                cur.execute("DELETE FROM global_recipient_locks WHERE recipient=? AND lock_token=?", (recipient, lock_token))
            
            conn.commit()
        except Exception:
            conn.rollback()
            raise

def release_send_lock(lead_id, lock_token, recipient, activity_id=None, status="Ready to Contact"):
    with get_db() as conn:
        try:
            conn.execute("BEGIN IMMEDIATE")
            now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            conn.execute("""
                UPDATE leads 
                SET status=?, status_updated_at=?, lock_token=NULL 
                WHERE lead_id=? AND lock_token=? AND status='Sending...'
            """, (status, now, lead_id, lock_token))
            conn.execute("DELETE FROM global_recipient_locks WHERE recipient=? AND lock_token=?", (recipient, lock_token))
            
            if activity_id:
                conn.execute("UPDATE activities SET status='Aborted', error='Pre-flight validation failed' WHERE activity_id=? AND lock_token=?", (activity_id, lock_token))
            conn.commit()
        except Exception:
            conn.rollback()
            raise

def reset_stuck_sends():
    with get_db() as conn:
        try:
            conn.execute("BEGIN IMMEDIATE")
            time_limit = (datetime.datetime.now() - datetime.timedelta(minutes=15)).strftime("%Y-%m-%d %H:%M:%S")
            cur = conn.cursor()
            
            cur.execute("SELECT lead_id, lock_token FROM leads WHERE status='Sending...' AND status_updated_at < ? AND lock_token IS NOT NULL", (time_limit,))
            stuck = cur.fetchall()
            
            if not stuck: 
                conn.commit()
                return 0, 0
            
            tokens = [t[1] for t in stuck]
            placeholders = ','.join(['?'] * len(tokens))
            
            cur.execute(f"UPDATE activities SET status='Unknown', error='Recovered after stale send timeout' WHERE status='Attempting' AND lock_token IN ({placeholders}) AND attempted_at < ?", tuple(tokens) + (time_limit,))
            act_count = cur.rowcount
            
            cur.execute(f"DELETE FROM global_recipient_locks WHERE lock_token IN ({placeholders})", tuple(tokens))
            
            cur.execute(f"UPDATE leads SET status='Send outcome unknown', status_updated_at=datetime('now','localtime'), lock_token=NULL WHERE status='Sending...' AND lock_token IN ({placeholders}) AND status_updated_at < ?", tuple(tokens) + (time_limit,))
            lead_count = cur.rowcount
                
            conn.commit()
            return lead_count, act_count
        except Exception:
            conn.rollback()
            raise




# ============================================================================
# EMAIL AUTOMATION (SMTP)
# ============================================================================


def fill_template(text, company_name):
    """[Company Name] / {{company}} → the prospect's name."""
    name = str(company_name or "").strip()
    return str(text or "").replace("[Company Name]", name).replace("{{company}}", name)


def connect_smtp(sender_email, sender_password):
    # Hard 60-second SMTP timeout keeps the send operation well below 
    # the stale-lock recovery window under normal execution.
    server = smtplib.SMTP("smtp.gmail.com", 587, timeout=60)
    server.starttls()
    server.login(sender_email, sender_password.replace(" ", ""))
    return server

def send_single_email(srv, sender_email, sender_password, receiver, email_subject, company_name, html_template):
    personalized_html = fill_template(html_template, company_name)
    soup_plain = BeautifulSoup(personalized_html, "html.parser")
    for tag in soup_plain(["style", "script"]): tag.decompose()
    
    msg = MIMEMultipart("alternative")
    msg["From"] = sender_email
    msg["To"] = receiver
    msg["Subject"] = email_subject
    msg.attach(MIMEText(soup_plain.get_text(separator="\n").strip(), "plain"))
    msg.attach(MIMEText(personalized_html, "html"))
    
    try:
        srv.send_message(msg)
        return True, "Contacted", srv
    except smtplib.SMTPRecipientsRefused: 
        return False, "Recipient rejected", srv
    except smtplib.SMTPSenderRefused: 
        return False, "Sender rejected", srv
    except Exception: 
        # If we cannot prove it was explicitly rejected before acceptance, treat as unknown
        return False, "Send outcome unknown", srv


# ============================================================================
# ICP GATING, SCORING & CONTEXTUAL NEGATION
# ============================================================================


def check_keyword_match(text, aliases):
    negated_found = None
    sentences = re.split(r'[.!?;]', text)
    negation_words = {"not", "no", "don't", "never", "without", "sin", "ningun", "ningún", "ninguna"}
    
    for sentence in sentences:
        for alias in aliases:
            for match in re.finditer(r'\b' + re.escape(alias) + r'\b', sentence):
                tokens_before = re.findall(r"\b[\w'-]+\b", sentence[:match.start()].lower())
                window_tokens = set(tokens_before[-5:])
                if not window_tokens.intersection(negation_words):
                    return True, alias, False
                else:
                    negated_found = alias
    if negated_found: return False, negated_found, True
    return False, None, False

def get_dual_score(biz_text, email, phone, whatsapp, wa_link_direct, website, req_kws, strong_kws, bonus_kws, exclude_kws):
    for kw_group in exclude_kws:
        matched, alias, negated = check_keyword_match(biz_text, kw_group)
        if matched: return 0, 0, 0, "❌ EXCLUDED", f"[EXC] {alias}", "EXCLUDED", True

    matched_words = []
    if req_kws:
        for req_group in req_kws:
            matched, alias, negated = check_keyword_match(biz_text, req_group)
            if matched: matched_words.append(f"[REQ] {alias}")
            else:
                if negated: matched_words.append(f"[NEGATED REQ] {alias}")
                return 0, 0, 0, "❌ NOT ICP", "Failed Required Gate", "NOT_ICP", False

    contactability_score = 0
    if email != "Not available": contactability_score += 20
    if phone != "Not available" or whatsapp != "Not available" or wa_link_direct != "Not available": contactability_score += 10
    if website != "Not available": contactability_score += 10

    icp_fit_raw = 0
    for kw_dict in strong_kws:
        matched, alias, negated = check_keyword_match(biz_text, kw_dict['aliases'])
        if matched:
            icp_fit_raw += kw_dict['weight']
            matched_words.append(f"[STR] {alias}")
        elif negated: matched_words.append(f"[NEGATED STR] {alias}")

    for kw_dict in bonus_kws:
        matched, alias, negated = check_keyword_match(biz_text, kw_dict['aliases'])
        if matched:
            icp_fit_raw += kw_dict['weight']
            matched_words.append(f"[BON] {alias}")
        elif negated: matched_words.append(f"[NEGATED BON] {alias}")

    icp_fit_score = min(60, icp_fit_raw)
    priority_score = contactability_score + icp_fit_score
    evidence_str = ", ".join(matched_words) if matched_words else "None"
    
    if priority_score >= 80: badge = "🔥 EXCELLENT"
    elif priority_score >= 50: badge = "🟢 GOOD"
    elif priority_score >= 30: badge = "🟡 NORMAL"
    else: badge = "🔴 LOW"
        
    return contactability_score, icp_fit_score, priority_score, badge, evidence_str, "QUALIFIED", False



# ============================================================================
# 6. INTERFACE TEXTS (English / Español)
# ============================================================================
TXT = {
"en": {
    "subtitle": "Developed by Código Origami · Alejandro Moreno | Pipeline Architecture V{v}",
    "lang_label": "Language",
    "tab1": "🔎 1. DISCOVER", "tab2": "🧠 2. QUALIFY", "tab3": "🎯 3. PROSPECT",
    # Tab 1
    "icp_title": "🎯 Define your Ideal Customer Profile (ICP)",
    "disc_params": "Discovery Parameters",
    "locations": "Locations (City, Country):",
    "locations_help": "Where to search, one per line, as 'City, Country' (e.g. 'Santiago, Chile'). The country is also used to read phone numbers and WhatsApp correctly. Google Maps returns at most ~120 places per search, so to get more results add several cities of the same country.",
    "queries": "Target Queries (Maps Search):",
    "queries_help": "What to type in Google Maps, one per line (e.g. 'auto parts importer'). Every query is searched in every location: 3 queries × 4 cities = 12 searches.",
    "max_results": "Max prospects per query:",
    "max_results_help": "Maximum number of places taken from each search before enrichment. Google Maps rarely shows more than ~120 per search; higher values only matter for very dense cities.",
    "mode": "Enrichment Mode:",
    "mode_fast": "⚡ FAST (Maps + Homepage)", "mode_smart": "🧠 SMART (Deepen only when promising or uncertain)",
    "mode_deep": "🔬 DEEP (Always full enrichment)",
    "mode_help": "FAST: Google Maps data + the company's homepage (quickest). SMART (recommended): goes deeper (contact pages, JavaScript websites, Facebook, search engine) only for prospects that pass the ICP gate and look promising or still have no reliable email. DEEP: does everything for every prospect (slowest, most complete).",
    "speed_box": "⚙️ Speed & reliability",
    "tabs": "Parallel Google Maps tabs",
    "tabs_help": "How many company pages are read at the same time. 2 is a safe default; 3–4 is faster but Google may show a CAPTCHA sooner. Use 1 for very large runs.",
    "fast_browser": "Fast browser (don't load images)",
    "fast_browser_help": "Skips images, videos and fonts while browsing. Much faster and uses less data; the information extracted is the same.",
    "show_browser": "Show the browser window",
    "show_browser_help": "Shows Chrome while it works. Useful to watch the process and to solve a CAPTCHA by hand if Google asks for one (the run then continues by itself).",
    "verify_mx": "Verify email domains (MX)",
    "verify_mx_help": "Checks that the email's domain can really receive mail before saving it. Avoids bounces that damage your sender reputation. Needs the 'dnspython' library.",
    "search_engine": "Search engine fallback",
    "search_engine_help": "If no reliable email is found on the website or social media, searches the company on DuckDuckGo. Only emails clearly tied to the company are kept. Slower; automatically paused if DuckDuckGo starts blocking.",
    "social_search": "Find social profiles when the website has none",
    "social_search_help": "Searches the company's Facebook / Instagram / LinkedIn on the web when Google Maps and the website don't link to them. Facebook pages often show an email.",
    "refresh_days": "Re-enrich known prospects older than (days)",
    "refresh_days_help": "Prospects already in your database are skipped (much faster) unless their data is older than this or you changed the ICP rules. Contacted / replied / excluded prospects are never overwritten.",
    "debug": "🐞 Debug mode",
    "debug_help": "Shows technical errors in the live log and keeps them in the 'debug' column. Only needed to diagnose problems.",
    "icp_box": "ICP Qualification (Gate & Weighted Signals)",
    "req": "Required GATE:",
    "req_help": "Words the prospect MUST have (in its Maps category, name or website). Each line is an AND condition; words separated by '|' are alternatives (OR). Prospects that fail are marked NOT ICP.",
    "req_caption": "Each line acts as an AND condition. Words divided by '|' act as OR.",
    "strong": "Strong Signals (kw | alias : weight 1-30):",
    "strong_help": "Words that make a prospect clearly interesting, with their weight. Example 'japanese | japan : 25' adds 25 points if 'japanese' or 'japan' appears. The ICP score is capped at 60.",
    "bonus": "Bonus Signals (kw : weight):",
    "bonus_help": "Secondary words that add a few points (e.g. 'b2b : 10').",
    "excl": "Absolute Exclusion (❌ EXCLUDED):",
    "excl_help": "If any of these words appears, the prospect is EXCLUDED and never contacted (e.g. 'rental, repair | mechanic'). Negations are understood: 'we do not do rentals' does not exclude.",
    "estimate": "⏳ Estimated time: ~{m} min · {s} searches",
    "choose": "Choose options",
    "start": "🚀 INITIATE PROSPECTING ENGINE",
    "start_help": "Starts searching Google Maps and enriching every prospect. You can follow the progress below; results are saved as they arrive.",
    "need_input": "❌ Please provide at least one Location and one Target Query.",
    "running": "Prospecting engine running…",
    "live_counters": "**Found:** {found} · **Processed:** {done} · **New:** {new} · **Updated:** {updated} · **Already known:** {known} · **With email:** {emails} · **ICP qualified:** {qualified} · **Excluded:** {excluded}",
    "done": "🎉 BATCH COMPLETE in {m} min.",
    "aborted": "⛔ Google blocked the connection, so the run stopped early. Everything found so far is saved. Try again in 1–2 hours or from another network.",
    "warn_no_identity": "⚠️ {n} prospects were dropped because they had no identifiable data (no Maps ID or address).",
    "warn_conflicts": "⚠️ {n} identity conflicts were logged for manual review.",
    "warn_blocked_q": "⚠️ {n} searches were blocked by Google; run them again later.",
    "log_searching": "🔎 '{q}' in {loc}", "log_results": "   {n} places found",
    "log_query_blocked": "🛑 Blocked by Google: '{q}' in {loc}",
    "log_unreadable": "⚠️ Could not read: {url}",
    "log_timeout": "⏱️ Enrichment timeout, saved without email: {name}",
    "log_blocked_pause": "🛑 Google is rate-limiting this connection ({n}/{m}). All tabs paused {mins} min. If the browser is visible you can solve the CAPTCHA to continue sooner.",
    "log_blocked_stop": "⛔ Google keeps blocking this connection. Stopping; progress is saved.",
    "log_captcha_ok": "✅ CAPTCHA solved, resuming.",
    "log_ddg_blocked": "⚠️ DuckDuckGo is blocking searches: search-engine fallback paused for this run.",
    # Tab 2
    "dash_title": "🧠 Prospect Qualification & Funnel Dashboard",
    "m_total": "TOTAL PROSPECTS", "m_total_help": "Every company in your local database.",
    "m_qual": "ICP QUALIFIED", "m_qual_help": "Passed the required gate and no exclusion word.",
    "m_noticp": "NOT ICP", "m_noticp_help": "Failed the required gate.",
    "m_excl": "EXCLUDED", "m_excl_help": "Contain an exclusion word: never contacted.",
    "m_email": "EMAIL READY", "m_email_help": "ICP qualified and with a verified email.",
    "m_wa": "WHATSAPP AVAILABLE", "m_wa_help": "With a WhatsApp link or a mobile number.",
    "m_contacted": "CONTACTED", "m_contacted_help": "Emailed from tab 3.",
    "m_replied": "REPLIED", "m_replied_help": "Marked as replied.",
    "filters": "Database Filters",
    "f_icp": "ICP Status:", "f_icp_help": "Show only prospects with these qualification results.",
    "f_badge": "Priority Badge:", "f_badge_help": "🔥 80+ points · 🟢 50–79 · 🟡 30–49 · 🔴 under 30 (contactability + ICP fit).",
    "f_email": "☑ Must have Email", "f_email_help": "Hide prospects without an email.",
    "f_phone": "☑ Must have Phone or WhatsApp", "f_phone_help": "Hide prospects without any phone, WhatsApp number or WhatsApp link.",
    "f_min": "Minimum Priority Score (/100)", "f_min_help": "Hide prospects below this total score.",
    "f_loc": "Location:", "f_loc_help": "Show only these locations (empty = all).",
    "f_search": "Search", "f_search_help": "Find text in company name, category, email, website or keywords.",
    "f_search_ph": "company, email, category…",
    "showing": "**Showing {n} prospects matching current filters.**",
    "excel": "📊 Export to Excel (formatted)", "excel_help": "Downloads the filtered table as an Excel file that looks like the screen: same columns, colours per priority, clickable links, filters and a summary sheet.",
    "csv": "💾 Export CSV", "csv_help": "Plain data (semicolon-separated, UTF-8) for other tools.",
    "recover": "🚨 Recover stuck sends", "recover_help": "If a campaign crashed or was closed mid-send, unlocks prospects stuck in 'Sending...' for more than 15 minutes (marked 'Send outcome unknown' for manual review).",
    "recovered": "Recovered {a} leads and {b} activities. Manual review required.",
    "danger": "🗑️ Danger zone", "confirm_clear": "I understand this deletes every prospect and campaign history",
    "confirm_clear_help": "Deleting the database cannot be undone. Export to Excel first if you need the data.",
    "clear": "🗑️ Clear Local Database", "clear_help": "Deletes the whole local database to start a new project.",
    "clear_blocked": "Cannot clear the database while sends are active. Wait for the campaign to finish or recover stuck sends.",
    "empty_db": "Your local database is empty. Run the Discovery engine in tab 1.",
    # Tab 3
    "out_title": "✉️ Controlled Outreach Campaign",
    "out_sub": "Send targeted cold emails safely, with global recipient locks: nobody is emailed twice.",
    "sender": "Sender Email Address", "sender_help": "Your Gmail or Google Workspace address. Emails are sent from it via Gmail's SMTP server.",
    "password": "App Password", "password_help": "NOT your normal password: a 16-letter App Password created in your Google account (see the guide below). It is used only during this session and never stored.",
    "subject": "Email Subject", "subject_help": "Subject line of the email. You can use [Company Name] to personalise it.",
    "template": "Email template", "template_help": "HTML files named template_*.html in the app folder (e.g. template_en.html, template_es.html). [Company Name] is replaced with each prospect's name.",
    "no_templates": "No template_*.html file found in the app folder.",
    "t_quality": "Send ONLY to leads with Priority Badge:", "t_quality_help": "Only ICP-qualified prospects with a valid email and these badges receive the email, best scores first.",
    "max_emails": "Max successful emails per campaign (Limit: 200):", "max_emails_help": "Stops after this many successful sends. For cold email, 30–50 a day per account protects your reputation.",
    "delay": "Humanized Delay between emails (seconds):", "delay_help": "Random wait between two emails so the sending pattern looks human and Gmail doesn't flag the account.",
    "preview": "👁️ Preview the email", "preview_company": "Sample Company Ltd",
    "test": "🧪 Send me a test email", "test_help": "Sends the selected template to your own address, so you can check it in your inbox.",
    "test_ok": "✅ Test sent to {e}.", "need_creds": "Please provide your sender email, App Password and subject.",
    "smtp_fail": "SMTP login failed: {e}. Check the App Password (see the guide below).",
    "start_camp": "🚀 START OUTREACH CAMPAIGN", "start_camp_help": "Sends the emails now, one by one with the delay above. Keep this tab open until it finishes.",
    "camp_log": "#### Campaign Log:",
    "sent_line": "[{i}/{n}] Sent to {c} ({e}) · waiting {s}s",
    "fail_line": "Failed to send to {c}: {r}",
    "camp_stop": "⚠️ Campaign stopped because of a sending/infrastructure issue. Sent {n} emails. Database updated.",
    "camp_early": "ℹ️ No more eligible prospects. Sent {n} emails. Database updated.",
    "camp_done": "✅ Campaign completed! Sent {n} emails. Database updated.",
    "no_lock": "No additional eligible leads could be locked in this batch.",
    "guide": "📧 How to create a Gmail App Password",
    "guide_text": "1. Open https://myaccount.google.com → **Security**.\n2. Turn on **2-Step Verification** (required).\n3. Search for **App passwords** and create one (e.g. 'Origami Prospector').\n4. Copy the 16-letter code without spaces and paste it in **App Password**.",
},
"es": {
    "subtitle": "Desarrollado por Código Origami · Alejandro Moreno | Arquitectura V{v}",
    "lang_label": "Idioma",
    "tab1": "🔎 1. DESCUBRIR", "tab2": "🧠 2. CUALIFICAR", "tab3": "🎯 3. PROSPECTAR",
    "icp_title": "🎯 Define tu Perfil de Cliente Ideal (ICP)",
    "disc_params": "Parámetros de búsqueda",
    "locations": "Ubicaciones (Ciudad, País):",
    "locations_help": "Dónde buscar, una por línea, como 'Ciudad, País' (p. ej. 'Santiago, Chile'). El país también sirve para leer bien los teléfonos y el WhatsApp. Google Maps devuelve como máximo ~120 sitios por búsqueda: para conseguir más, añade varias ciudades del mismo país.",
    "queries": "Búsquedas objetivo (Google Maps):",
    "queries_help": "Lo que se escribe en Google Maps, una por línea (p. ej. 'importador de recambios'). Cada búsqueda se hace en cada ubicación: 3 búsquedas × 4 ciudades = 12 búsquedas.",
    "max_results": "Máximo de prospectos por búsqueda:",
    "max_results_help": "Número máximo de sitios que se toman de cada búsqueda antes de enriquecerlos. Google Maps rara vez muestra más de ~120 por búsqueda; valores mayores solo importan en ciudades muy grandes.",
    "mode": "Modo de enriquecimiento:",
    "mode_fast": "⚡ RÁPIDO (Maps + página principal)", "mode_smart": "🧠 INTELIGENTE (profundiza solo si promete o hay dudas)",
    "mode_deep": "🔬 PROFUNDO (siempre completo)",
    "mode_help": "RÁPIDO: datos de Google Maps + la página principal de la web (lo más veloz). INTELIGENTE (recomendado): profundiza (páginas de contacto, webs con JavaScript, Facebook, buscador) solo en los prospectos que pasan el filtro ICP y prometen o aún no tienen un email fiable. PROFUNDO: lo hace todo con todos (más lento y más completo).",
    "speed_box": "⚙️ Velocidad y fiabilidad",
    "tabs": "Pestañas de Google Maps en paralelo",
    "tabs_help": "Cuántas fichas de empresa se leen a la vez. 2 es un valor seguro; 3–4 es más rápido, pero Google puede mostrar un CAPTCHA antes. Usa 1 para búsquedas muy grandes.",
    "fast_browser": "Navegador rápido (sin imágenes)",
    "fast_browser_help": "No carga imágenes, vídeos ni fuentes al navegar. Mucho más rápido y gasta menos datos; la información extraída es la misma.",
    "show_browser": "Mostrar la ventana del navegador",
    "show_browser_help": "Muestra Chrome mientras trabaja. Útil para ver el proceso y para resolver a mano un CAPTCHA si Google lo pide (después continúa solo).",
    "verify_mx": "Verificar dominios de email (MX)",
    "verify_mx_help": "Comprueba que el dominio del email puede recibir correo antes de guardarlo. Evita rebotes que dañan la reputación de tu cuenta. Necesita la librería 'dnspython'.",
    "search_engine": "Buscar en buscador si falta email",
    "search_engine_help": "Si no se encuentra un email fiable en la web o en redes sociales, busca la empresa en DuckDuckGo. Solo se guardan emails claramente ligados a la empresa. Más lento; se pausa solo si DuckDuckGo empieza a bloquear.",
    "social_search": "Buscar redes sociales si la web no las tiene",
    "social_search_help": "Busca el Facebook / Instagram / LinkedIn de la empresa en la web cuando ni Google Maps ni su web los enlazan. Las páginas de Facebook suelen mostrar un email.",
    "refresh_days": "Volver a enriquecer prospectos conocidos con más de (días)",
    "refresh_days_help": "Los prospectos que ya están en tu base de datos se saltan (mucho más rápido) salvo que sus datos tengan más días que este valor o hayas cambiado las reglas ICP. Los contactados / que respondieron / excluidos nunca se sobrescriben.",
    "debug": "🐞 Modo diagnóstico",
    "debug_help": "Muestra los errores técnicos en el registro y los guarda en la columna 'debug'. Solo hace falta para diagnosticar problemas.",
    "icp_box": "Cualificación ICP (filtro y señales ponderadas)",
    "req": "Filtro OBLIGATORIO:",
    "req_help": "Palabras que el prospecto DEBE tener (en su categoría de Maps, nombre o web). Cada línea es una condición Y; las palabras separadas por '|' son alternativas (O). Los que no cumplen quedan como NO ICP.",
    "req_caption": "Cada línea es una condición Y. Las palabras separadas por '|' funcionan como O.",
    "strong": "Señales fuertes (palabra | alias : peso 1-30):",
    "strong_help": "Palabras que hacen claramente interesante a un prospecto, con su peso. Ejemplo 'japanese | japan : 25' suma 25 puntos si aparece 'japanese' o 'japan'. La puntuación ICP tiene un máximo de 60.",
    "bonus": "Señales extra (palabra : peso):",
    "bonus_help": "Palabras secundarias que suman algunos puntos (p. ej. 'b2b : 10').",
    "excl": "Exclusión absoluta (❌ EXCLUIDO):",
    "excl_help": "Si aparece alguna de estas palabras, el prospecto queda EXCLUIDO y nunca se contacta (p. ej. 'rental, repair | mechanic'). Entiende las negaciones: 'no hacemos alquileres' no excluye.",
    "estimate": "⏳ Tiempo estimado: ~{m} min · {s} búsquedas",
    "choose": "Elige opciones",
    "start": "🚀 INICIAR MOTOR DE PROSPECCIÓN",
    "start_help": "Empieza a buscar en Google Maps y a enriquecer cada prospecto. Puedes seguir el progreso aquí abajo; los resultados se guardan a medida que llegan.",
    "need_input": "❌ Indica al menos una ubicación y una búsqueda.",
    "running": "Motor de prospección en marcha…",
    "live_counters": "**Encontrados:** {found} · **Procesados:** {done} · **Nuevos:** {new} · **Actualizados:** {updated} · **Ya conocidos:** {known} · **Con email:** {emails} · **ICP cualificados:** {qualified} · **Excluidos:** {excluded}",
    "done": "🎉 ¡LOTE COMPLETADO en {m} min!",
    "aborted": "⛔ Google bloqueó la conexión y la ejecución se detuvo antes de tiempo. Todo lo encontrado está guardado. Vuelve a intentarlo en 1–2 horas o desde otra red.",
    "warn_no_identity": "⚠️ Se descartaron {n} prospectos sin datos identificables (sin ID de Maps ni dirección).",
    "warn_conflicts": "⚠️ Se registraron {n} conflictos de identidad para revisión manual.",
    "warn_blocked_q": "⚠️ Google bloqueó {n} búsquedas; vuelve a lanzarlas más tarde.",
    "log_searching": "🔎 '{q}' en {loc}", "log_results": "   {n} sitios encontrados",
    "log_query_blocked": "🛑 Bloqueado por Google: '{q}' en {loc}",
    "log_unreadable": "⚠️ No se pudo leer: {url}",
    "log_timeout": "⏱️ Tiempo agotado al enriquecer, guardado sin email: {name}",
    "log_blocked_pause": "🛑 Google está frenando esta conexión ({n}/{m}). Todas las pestañas en pausa {mins} min. Si el navegador está visible, puedes resolver el CAPTCHA para seguir antes.",
    "log_blocked_stop": "⛔ Google sigue bloqueando esta conexión. Se detiene; el progreso está guardado.",
    "log_captcha_ok": "✅ CAPTCHA resuelto, continuando.",
    "log_ddg_blocked": "⚠️ DuckDuckGo está bloqueando búsquedas: búsqueda en buscador pausada en esta ejecución.",
    "dash_title": "🧠 Cualificación de prospectos y embudo",
    "m_total": "PROSPECTOS TOTALES", "m_total_help": "Todas las empresas de tu base de datos local.",
    "m_qual": "ICP CUALIFICADOS", "m_qual_help": "Pasan el filtro obligatorio y no tienen palabras de exclusión.",
    "m_noticp": "NO ICP", "m_noticp_help": "No pasan el filtro obligatorio.",
    "m_excl": "EXCLUIDOS", "m_excl_help": "Contienen una palabra de exclusión: nunca se contactan.",
    "m_email": "CON EMAIL", "m_email_help": "ICP cualificados y con email verificado.",
    "m_wa": "CON WHATSAPP", "m_wa_help": "Con enlace de WhatsApp o número móvil.",
    "m_contacted": "CONTACTADOS", "m_contacted_help": "Recibieron email desde la pestaña 3.",
    "m_replied": "RESPONDIERON", "m_replied_help": "Marcados como 'respondió'.",
    "filters": "Filtros de la base de datos",
    "f_icp": "Estado ICP:", "f_icp_help": "Muestra solo prospectos con estos resultados de cualificación.",
    "f_badge": "Prioridad:", "f_badge_help": "🔥 80+ puntos · 🟢 50–79 · 🟡 30–49 · 🔴 menos de 30 (contactabilidad + encaje ICP).",
    "f_email": "☑ Con email obligatorio", "f_email_help": "Oculta los prospectos sin email.",
    "f_phone": "☑ Con teléfono o WhatsApp", "f_phone_help": "Oculta los prospectos sin teléfono, número de WhatsApp ni enlace de WhatsApp.",
    "f_min": "Puntuación mínima (/100)", "f_min_help": "Oculta los prospectos por debajo de esta puntuación total.",
    "f_loc": "Ubicación:", "f_loc_help": "Muestra solo estas ubicaciones (vacío = todas).",
    "f_search": "Buscar", "f_search_help": "Busca texto en el nombre, categoría, email, web o palabras clave.",
    "f_search_ph": "empresa, email, categoría…",
    "showing": "**Mostrando {n} prospectos con los filtros actuales.**",
    "excel": "📊 Exportar a Excel (con formato)", "excel_help": "Descarga la tabla filtrada como un Excel igual que en pantalla: mismas columnas, colores por prioridad, enlaces clicables, filtros y una hoja de resumen.",
    "csv": "💾 Exportar CSV", "csv_help": "Datos sin formato (separados por punto y coma, UTF-8) para otras herramientas.",
    "recover": "🚨 Recuperar envíos atascados", "recover_help": "Si una campaña falló o se cerró a mitad de envío, desbloquea los prospectos atascados en 'Sending...' más de 15 minutos (quedan como 'Send outcome unknown' para revisarlos a mano).",
    "recovered": "Recuperados {a} prospectos y {b} actividades. Requiere revisión manual.",
    "danger": "🗑️ Zona peligrosa", "confirm_clear": "Entiendo que esto borra todos los prospectos y el historial de campañas",
    "confirm_clear_help": "Borrar la base de datos no se puede deshacer. Exporta antes a Excel si necesitas los datos.",
    "clear": "🗑️ Borrar base de datos local", "clear_help": "Borra toda la base de datos local para empezar un proyecto nuevo.",
    "clear_blocked": "No se puede borrar mientras hay envíos activos. Espera a que termine la campaña o recupera los envíos atascados.",
    "empty_db": "Tu base de datos local está vacía. Lanza el motor de búsqueda en la pestaña 1.",
    "out_title": "✉️ Campaña de emails controlada",
    "out_sub": "Envía emails en frío de forma segura, con bloqueos globales por destinatario: nadie recibe el email dos veces.",
    "sender": "Email remitente", "sender_help": "Tu dirección de Gmail o Google Workspace. Los emails se envían desde ella con el servidor SMTP de Gmail.",
    "password": "Contraseña de aplicación", "password_help": "NO es tu contraseña normal: es una contraseña de aplicación de 16 letras creada en tu cuenta de Google (ver la guía de abajo). Solo se usa en esta sesión y nunca se guarda.",
    "subject": "Asunto del email", "subject_help": "Asunto del email. Puedes usar [Company Name] para personalizarlo.",
    "template": "Plantilla de email", "template_help": "Archivos HTML llamados template_*.html en la carpeta de la app (p. ej. template_en.html, template_es.html). [Company Name] se sustituye por el nombre de cada prospecto.",
    "no_templates": "No hay ningún archivo template_*.html en la carpeta de la app.",
    "t_quality": "Enviar SOLO a prospectos con prioridad:", "t_quality_help": "Solo reciben el email los prospectos ICP cualificados con email válido y estas prioridades, empezando por las mejores puntuaciones.",
    "max_emails": "Máximo de emails enviados por campaña (límite: 200):", "max_emails_help": "Se detiene tras este número de envíos correctos. En emails en frío, 30–50 al día por cuenta protegen tu reputación.",
    "delay": "Pausa humanizada entre emails (segundos):", "delay_help": "Espera aleatoria entre dos emails para que el ritmo parezca humano y Gmail no marque la cuenta.",
    "preview": "👁️ Vista previa del email", "preview_company": "Empresa de Ejemplo S.L.",
    "test": "🧪 Enviarme un email de prueba", "test_help": "Envía la plantilla elegida a tu propia dirección para que la revises en tu bandeja.",
    "test_ok": "✅ Prueba enviada a {e}.", "need_creds": "Indica el email remitente, la contraseña de aplicación y el asunto.",
    "smtp_fail": "Error al iniciar sesión SMTP: {e}. Revisa la contraseña de aplicación (ver la guía de abajo).",
    "start_camp": "🚀 INICIAR CAMPAÑA DE EMAILS", "start_camp_help": "Envía los emails ahora, uno a uno con la pausa indicada. Mantén esta pestaña abierta hasta que termine.",
    "camp_log": "#### Registro de la campaña:",
    "sent_line": "[{i}/{n}] Enviado a {c} ({e}) · esperando {s}s",
    "fail_line": "No se pudo enviar a {c}: {r}",
    "camp_stop": "⚠️ Campaña detenida por un problema de envío/infraestructura. Enviados {n} emails. Base de datos actualizada.",
    "camp_early": "ℹ️ No quedan prospectos elegibles. Enviados {n} emails. Base de datos actualizada.",
    "camp_done": "✅ ¡Campaña completada! Enviados {n} emails. Base de datos actualizada.",
    "no_lock": "No se pudieron bloquear más prospectos elegibles en este lote.",
    "guide": "📧 Cómo crear una contraseña de aplicación de Gmail",
    "guide_text": "1. Abre https://myaccount.google.com → **Seguridad**.\n2. Activa la **Verificación en dos pasos** (obligatoria).\n3. Busca **Contraseñas de aplicaciones** y crea una (p. ej. 'Origami Prospector').\n4. Copia el código de 16 letras sin espacios y pégalo en **Contraseña de aplicación**.",
},
}

COLUMN_LABELS = {
    "en": {"company": "Company", "category": "Category", "location": "Location", "website": "Website",
           "maps_url": "Google Maps", "email": "Email", "email_confidence": "Email confidence",
           "email_source": "Email source", "phone": "Phone", "whatsapp": "WhatsApp", "wa_link": "WhatsApp link",
           "facebook": "Facebook", "instagram": "Instagram", "linkedin": "LinkedIn",
           "contactability_score": "Contactability", "icp_fit_score": "ICP fit", "priority_score": "Priority",
           "priority_badge": "Badge", "qual_status": "ICP status", "status": "Status",
           "matched_keywords": "Matched keywords", "enrichment_level": "Enrichment", "address": "Address",
           "enriched_at": "Updated", "last_contacted_at": "Last contacted", "sources_searched": "Sources searched",
           "debug_logs": "Debug"},
    "es": {"company": "Empresa", "category": "Categoría", "location": "Ubicación", "website": "Web",
           "maps_url": "Google Maps", "email": "Email", "email_confidence": "Fiabilidad email",
           "email_source": "Origen del email", "phone": "Teléfono", "whatsapp": "WhatsApp",
           "wa_link": "Enlace WhatsApp", "facebook": "Facebook", "instagram": "Instagram", "linkedin": "LinkedIn",
           "contactability_score": "Contactabilidad", "icp_fit_score": "Encaje ICP", "priority_score": "Prioridad",
           "priority_badge": "Nivel", "qual_status": "Estado ICP", "status": "Estado",
           "matched_keywords": "Palabras encontradas", "enrichment_level": "Enriquecimiento", "address": "Dirección",
           "enriched_at": "Actualizado", "last_contacted_at": "Último contacto", "sources_searched": "Fuentes consultadas",
           "debug_logs": "Debug"},
}


def t(key, lang="en", **kw):
    s = TXT.get(lang, TXT["en"]).get(key) or TXT["en"].get(key, key)
    return s.format(**kw) if kw else s




# ============================================================================
# 5. GOOGLE MAPS ENGINE (async Playwright: 1 search tab + N detail tabs in parallel)
# ============================================================================
PROFILE_DIR = os.path.join(BASE_DIR, "browser_profile")
READY_JS = """
() => !!document.querySelector('div[role="feed"]')
   || !!(document.querySelector('h1.DUwDvf') && document.querySelector('h1.DUwDvf').innerText.trim())
   || location.href.includes('/sorry/')
   || /can.t find|no results found|unusual traffic|not a robot/i.test((document.body && document.body.innerText) || '')
"""
PLACE_JS = """
() => {
  const q = s => document.querySelector(s);
  const main = q('div[role="main"]') || document.body;
  const h1 = q('h1.DUwDvf') || main.querySelector('h1');
  const cat = q('button[jsaction*="category"]') || q('button.DkEaL');
  const addr = q('button[data-item-id="address"]');
  const phone = q('button[data-item-id^="phone:tel:"]');
  const web = q('a[data-item-id="authority"]');
  return {
    name: h1 ? h1.innerText.trim() : '',
    category: cat ? cat.innerText.trim() : '',
    address: addr ? (addr.getAttribute('aria-label') || addr.innerText).replace(/^Address:\\s*/i, '').replace(/\\n/g, ', ').trim() : '',
    phone: phone ? phone.getAttribute('data-item-id').replace('phone:tel:', '') : '',
    website: web ? web.href : '',
    links: [...main.querySelectorAll('a[href]')].map(a => a.href),
    text: main.innerText || ''
  };
}
"""
BLOCK_RE = re.compile(r"unusual traffic|not a robot|tr[aá]fico inusual", re.I)


class RunUI:
    """Live progress inside Streamlit (metrics + the last log lines, cheap to redraw)."""
    def __init__(self, lang):
        self.lang = lang
        self.metrics = st.empty()
        self.box = st.empty()
        self.lines = []
        self.c = {"found": 0, "done": 0, "new": 0, "updated": 0, "known": 0, "emails": 0,
                  "qualified": 0, "excluded": 0}

    def log(self, msg):
        self.lines = (self.lines + [msg])[-14:]
        self.box.code("\n".join(self.lines), language=None)

    def refresh(self):
        c = self.c
        self.metrics.markdown(t("live_counters", self.lang, **c))


class Gate:
    """Shared pause for all tabs when Google rate-limits the connection; stops the run if it persists."""
    MAX_TRIPS = 4

    def __init__(self, ui):
        self.ui, self.until, self.pause, self.trips, self.aborted, self.page = ui, 0.0, 120, 0, False, None

    async def wait(self):
        while time.time() < self.until and not self.aborted:
            if self.page is not None and "/sorry/" not in self.page.url:
                try:
                    if not await self.page.get_by_text(BLOCK_RE).count():
                        self.until = 0
                        self.ui.log(t("log_captcha_ok", self.ui.lang)); break
                except Exception:
                    pass
            await asyncio.sleep(5)

    def trip(self, page):
        if time.time() < self.until or self.aborted: return
        self.trips += 1
        if self.trips >= self.MAX_TRIPS:
            self.aborted = True
            self.ui.log(t("log_blocked_stop", self.ui.lang)); return
        self.until, self.page = time.time() + self.pause, page
        self.ui.log(t("log_blocked_pause", self.ui.lang, n=self.trips, m=self.MAX_TRIPS - 1, mins=self.pause // 60))
        self.pause = min(self.pause * 2, 900)

    def ok(self):
        self.trips = 0
        self.pause = max(120, self.pause // 2)


async def block_heavy(route):
    if route.request.resource_type in ("image", "media", "font"): await route.abort()
    else: await route.continue_()


async def handle_consent(page):
    if "consent." not in page.url: return
    for label in ("Accept all", "Aceptar todo", "Reject all", "Rechazar todo", "Tout accepter",
                  "Alle akzeptieren", "Aceitar tudo", "Accetta tutto"):
        btn = page.get_by_role("button", name=label)
        try:
            if await btn.count():
                await btn.first.click(timeout=4000)
                await page.wait_for_load_state("domcontentloaded")
                return
        except Exception:
            continue


async def is_blocked(page):
    if "/sorry/" in page.url: return True
    try: return await page.get_by_text(BLOCK_RE).count() > 0
    except Exception: return False


async def open_browser(p, cfg):
    """Persistent profile (cookies & consent survive) + real Chrome when installed = fewer CAPTCHAs."""
    kwargs = dict(user_data_dir=PROFILE_DIR, headless=cfg["headless"], locale="en-US",
                  viewport={"width": 1280, "height": 900})
    try:
        context = await p.chromium.launch_persistent_context(channel="chrome", **kwargs)
    except Exception:
        context = await p.chromium.launch_persistent_context(**kwargs)
    if cfg["fast_browser"]:
        await context.route("**/*", block_heavy)
    return context


async def collect_query(ctx, page, query, max_results):
    """URLs of the places for a search, [] if Maps has none, None if Google blocked us."""
    url = add_hl_param(f"https://www.google.com/maps/search/{quote_plus(query)}")
    for _ in range(3):
        await ctx.gate.wait()
        if ctx.gate.aborted: return None
        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=45000)
            await handle_consent(page)
            await page.wait_for_function(READY_JS, timeout=25000)
        except Exception:
            pass
        if await is_blocked(page):
            ctx.gate.trip(page); continue
        break
    else:
        return None
    feed = page.locator('div[role="feed"]')
    if await feed.count() == 0:
        return [page.url] if "/maps/place/" in page.url else []
    end_marker = page.get_by_text(re.compile(r"reached the end of the list", re.I))
    prev, stalls = 0, 0
    for _ in range(max(6, min(80, math.ceil(max_results / 5) + 6))):
        try: await feed.evaluate("el => el.scrollTo(0, el.scrollHeight)")
        except Exception: break
        await page.wait_for_timeout(random.randint(1000, 1700))
        if await end_marker.count(): break
        n = await page.locator('a[href*="/maps/place/"]').count()
        if n >= max_results: break
        stalls = stalls + 1 if n == prev else 0
        if stalls >= 4: break
        prev = n
    ctx.gate.ok()
    hrefs = await page.eval_on_selector_all('a[href*="/maps/place/"]', "els => els.map(e => e.href)")
    urls = [clean_url(h) for h in hrefs if "/reviews/" not in h and "/photos/" not in h]
    return list(dict.fromkeys(urls))[:max_results]


async def scrape_place(ctx, page, url):
    for attempt in range(3):
        await ctx.gate.wait()
        if ctx.gate.aborted: return None
        await asyncio.sleep(random.uniform(0.5, 1.5))
        try:
            await page.goto(add_hl_param(url), wait_until="domcontentloaded", timeout=40000)
            await handle_consent(page)
            await page.wait_for_function(READY_JS, timeout=20000)
        except Exception:
            pass
        if await is_blocked(page):
            ctx.gate.trip(page); continue
        try:
            await page.wait_for_timeout(random.randint(500, 900))
            data = await page.evaluate(PLACE_JS)
            if data.get("name"):
                ctx.gate.ok()
                return data
        except Exception:
            pass
        if attempt: break
        await asyncio.sleep(3)
    return None


async def render_html(page, url, facebook=False):
    try:
        if facebook:
            url = (url.split("&sk=")[0] + "&sk=about") if "profile.php" in url else url.split("?")[0].rstrip("/") + "/about"
        await page.goto(url, wait_until="load", timeout=25000)
        await page.wait_for_timeout(2500 if facebook else 1500)
        if facebook:
            try:
                close = page.locator('[aria-label="Close"], [aria-label="Cerrar"]').first
                if await close.count(): await close.click(timeout=2000)
            except Exception:
                pass
        return (await page.content()) + "\n" + (await page.locator("body").inner_text(timeout=5000)), page.url
    except Exception:
        return None, url


async def with_page(pool, fn):
    page = await pool.get()
    try: return await fn(page)
    finally: pool.put_nowait(page)


async def run_in(ctx, fn, *args):
    return await asyncio.get_running_loop().run_in_executor(ctx.pool, fn, *args)


async def enrich(ctx, d, loc, region, icp):
    """Cascading enrichment. FAST: homepage. SMART: deeper only when promising or still uncertain.
    DEEP: everything. Returns the contact fields, ICP text and logs."""
    mode, logs, searched = ctx.cfg["mode"], [], ["Google Maps"]
    tokens = company_tokens(d["name"])
    website = clean_url(d.get("website") or "") or NA
    if website != NA and is_social(website):
        socials_seed = [website]; website_for_fetch = NA
    else:
        socials_seed = []; website_for_fetch = website
    site_domain = normalize_domain(website_for_fetch)
    cands = [(e, "Google Maps", d.get("maps_url", NA)) for e in emails_in_text(d.get("text", ""))]
    wa_found = wa_from_text(d.get("text", ""))
    tels, socials, signal = [], set(socials_seed) | set(x for x in d.get("links", []) if is_social(x)), ""

    def score_now():
        bt = f"{d.get('category', '')} {d['name']} {signal}".lower()
        has_email = NA if not cands else "x"
        return get_dual_score(bt, has_email, d.get("phone") or NA, NA, NA, website, *icp)

    acc = None
    if website_for_fetch != NA:
        searched.append("Website (Homepage)")
        acc = await run_in(ctx, site_static, website_for_fetch, mode == "DEEP")
        site_domain = acc["final_domain"] or site_domain
        cands += acc["emails"]; wa_found += acc["wa"]; tels += acc["tels"]; socials |= acc["socials"]
        signal = acc["signal"]; logs += acc["errors"]

    c, b, total, badge, ev, qual, excl = score_now()
    if mode == "DEEP": go_deep = True
    elif mode == "SMART": go_deep = qual == "QUALIFIED" and (b >= 30 or best_conf(cands, site_domain, tokens) < 3)
    else: go_deep = False

    if go_deep:
        if acc and mode == "SMART" and best_conf(cands, site_domain, tokens) < 3:
            searched.append("Website (Contact pages)")
            more = await run_in(ctx, site_static, website_for_fetch, True)
            cands += [x for x in more["emails"] if x not in cands]; wa_found += more["wa"]
            tels += more["tels"]; socials |= more["socials"]; signal += " " + more["signal"][:2000]
            acc = more
        if acc and acc["needs_render"] and best_conf(cands, site_domain, tokens) < 3:
            searched.append("Website (Browser render)")
            html, final = await with_page(ctx.render_pages, lambda pg: render_html(pg, acc["home_url"]))
            if html:
                info = await run_in(ctx, extract_from_html, html, final)
                cands += [(e, "Website", final) for e in info["emails"]]
                wa_found += info["wa"]; socials |= info["socials"]; signal += " " + info["signal"][:2000]
        soc = split_socials(socials)
        if soc["facebook"] == NA and best_conf(cands, site_domain, tokens) < 2 and ctx.cfg["social_search"]:
            searched.append("Social Search")
            found = await run_in(ctx, search_socials, d["name"], loc)
            for k, v in (found or {}).items():
                if v != NA and soc.get(k) == NA: soc[k] = v; socials.add(v)
        if soc["facebook"] != NA and best_conf(cands, site_domain, tokens) < 3:
            searched.append("Facebook")
            html, final = await with_page(ctx.render_pages, lambda pg: render_html(pg, soc["facebook"], True))
            if html:
                info = await run_in(ctx, extract_from_html, html, "")
                cands += [(e, "Social Media", soc["facebook"]) for e in info["emails"]]
                wa_found += info["wa"]
        if best_conf(cands, site_domain, tokens) < 2 and ctx.cfg["search_engine"]:
            searched.append("Search Engine")
            for e in await run_in(ctx, search_emails, d["name"], loc):
                cands.append((e, "Search Engine", "Search Result"))
        if ddg_blocked_now() and not getattr(ctx, "ddg_warned", False):
            ctx.ddg_warned = True
            ctx.ui.log(t("log_ddg_blocked", ctx.ui.lang))

    email, src, src_url, conf = await run_in(ctx, choose_best_email, cands, site_domain, tokens, ctx.cfg["verify_mx"])

    whatsapp, wa_link = NA, NA
    prio = lambda it: 0 if (it[1] or "").startswith("http") else (1 if it[1] == "WA" else 2)
    for num, link in sorted(wa_found, key=prio):
        e164 = to_e164(num, region)
        if e164:
            whatsapp = e164
            if link and link.startswith("http"): wa_link = link
            elif link == "WA": wa_link = f"https://wa.me/{e164.lstrip('+')}"
            break
    if whatsapp == NA:
        for num in [d.get("phone")] + tels:
            if num and is_mobile(num, region):
                whatsapp = to_e164(num, region); break
    soc = split_socials(socials)
    return {"email": email, "email_source": src, "email_source_url": src_url, "email_confidence": conf,
            "whatsapp": whatsapp, "wa_link": wa_link, "facebook": soc["facebook"],
            "instagram": soc["instagram"], "linkedin": soc["linkedin"], "signal": signal,
            "level": "DEEP" if go_deep else "FAST", "searched": searched, "logs": logs,
            "website": website}


def ddg_blocked_now():
    return _ddg_state["blocked"]


HARD_TERMINAL = {"Contacted", "Replied", "Not Interested", "Do Not Contact", "Excluded", "Sending...",
                 "Send outcome unknown"}


async def run_discovery(cfg, ui):
    """Whole discovery run. Returns (run_id, summary dict)."""
    icp = (cfg["req"], cfg["strong"], cfg["bonus"], cfg["excl"])
    icp_hash = get_icp_hash(*icp)
    run_id = f"RUN-{uuid.uuid4().hex[:12].upper()}"
    ctx = types.SimpleNamespace(cfg=cfg, ui=ui, gate=Gate(ui), pool=ThreadPoolExecutor(max_workers=12))
    summary = {"no_identity": 0, "conflicts": 0, "blocked_queries": 0, "empty_queries": 0}
    db_conn = get_db()
    tasks = []
    seen = set()

    async def process(url, base, loc, region):
        try:
            d = await with_page(ctx.detail_pages, lambda pg: scrape_place(ctx, pg, url))
            if not d:
                if not ctx.gate.aborted: ui.log(t("log_unreadable", ui.lang, url=url[:60]))
                return
            maps_id = extract_place_id(url)
            phone_fmt = format_phone(d.get("phone"), region) if d.get("phone") else NA
            website = clean_url(d.get("website") or "") or NA
            domain = normalize_domain(website)
            norm_name, norm_addr = normalize_company_name(d["name"]), normalize_address(d.get("address") or NA)
            aliases = generate_identity_aliases(maps_id, domain, normalize_phone(phone_fmt), norm_name, norm_addr)

            # 1. Already known and up to date? Skip the expensive part.
            try:
                db_conn.execute("BEGIN IMMEDIATE")
                lead_id, id_status, conflict_ids = resolve_lead_identity(db_conn, aliases)
                if not lead_id:
                    if id_status == "IDENTITY_CONFLICT":
                        log_identity_conflict(db_conn, run_id, aliases, conflict_ids, d["name"], maps_id, domain, phone_fmt, d.get("address"))
                        summary["conflicts"] += 1
                    else:
                        summary["no_identity"] += 1
                    db_conn.commit(); return
                res = db_conn.execute("SELECT status, enriched_at, icp_hash FROM leads WHERE lead_id=?", (lead_id,)).fetchone()
                db_conn.commit()
            except Exception as e:
                db_conn.rollback(); ui.log(f"⚠️ DB: {type(e).__name__}"); return
            if res:
                status_val, enriched_at, saved_hash = res
                refresh = not enriched_at
                if not refresh and status_val not in HARD_TERMINAL:
                    try:
                        age = (datetime.datetime.now() - datetime.datetime.strptime(enriched_at, "%Y-%m-%d %H:%M:%S")).days
                        refresh = age >= cfg["refresh_days"] or saved_hash != icp_hash
                    except Exception:
                        refresh = True
                if not refresh:
                    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    try:
                        db_conn.execute("BEGIN IMMEDIATE")
                        db_conn.execute("UPDATE leads SET last_seen_at=? WHERE lead_id=?", (now, lead_id))
                        db_conn.execute("INSERT OR IGNORE INTO lead_discoveries (run_id, lead_id, query, location, discovered_at) VALUES (?,?,?,?,?)",
                                        (run_id, lead_id, base, loc, now))
                        db_conn.commit()
                    except Exception:
                        db_conn.rollback()
                    ui.c["known"] += 1
                    return

            # 2. Enrich
            started = time.time()
            d["maps_url"] = url
            try:
                e = await asyncio.wait_for(enrich(ctx, d, loc, region, icp), timeout=150)
            except asyncio.TimeoutError:
                ui.log(t("log_timeout", ui.lang, name=d["name"][:40]))
                e = {"email": NA, "email_source": NA, "email_source_url": NA, "email_confidence": "NONE",
                     "whatsapp": to_e164(d.get("phone"), region) if is_mobile(d.get("phone"), region) else NA,
                     "wa_link": NA, "facebook": NA, "instagram": NA, "linkedin": NA, "signal": "",
                     "level": "FAST", "searched": ["Google Maps"], "logs": ["Enrichment timeout"], "website": website}
            biz = f"{d.get('category', '')} {d['name']} {e['signal']}".lower()
            c, b, total, badge, ev, qual, excl = get_dual_score(
                biz, e["email"], phone_fmt, e["whatsapp"], e["wa_link"], website, *icp)
            now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            found = []
            if e["email"] != NA: found.append("Email")
            if e["whatsapp"] != NA or e["wa_link"] != NA: found.append("WhatsApp")
            if any(e[k] != NA for k in ("facebook", "instagram", "linkedin")): found.append("Social")
            lead = {
                "lead_id": lead_id, "created_at": now, "last_seen_at": now, "enriched_at": now,
                "maps_id": maps_id, "maps_url": url, "domain": domain, "norm_name": norm_name,
                "norm_address": norm_addr, "location": normalize_location(loc), "query": base,
                "company": d["name"], "category": d.get("category") or NA, "address": d.get("address") or NA,
                "website": website, "phone": phone_fmt, "whatsapp": e["whatsapp"], "email": e["email"],
                "email_source": e["email_source"], "email_source_url": e["email_source_url"],
                "email_confidence": e["email_confidence"], "wa_link": e["wa_link"],
                "facebook": e["facebook"], "instagram": e["instagram"], "linkedin": e["linkedin"],
                "contactability_score": c, "icp_fit_score": b, "priority_score": total,
                "matched_keywords": ev, "qual_status": qual, "priority_badge": badge,
                "status": "Excluded" if qual == "EXCLUDED" else "Not ICP" if qual == "NOT_ICP" else "Ready to Contact",
                "enrichment_level": e["level"], "enrich_time": f"{round(time.time() - started, 1)}s",
                "icp_hash": icp_hash, "sources_searched": ", ".join(e["searched"]),
                "sources_found": ", ".join(found) if found else "None",
                "debug_logs": "; ".join(e["logs"]) if e["logs"] else "None", "country": region or "",
            }
            disc = {"run_id": run_id, "query": base, "location": loc, "discovered_at": now}
            try:
                upsert_lead_transaction(db_conn, lead, aliases, disc)
            except Exception as ex:
                ui.log(f"⚠️ DB: {type(ex).__name__}: {str(ex)[:60]}"); return
            ui.c["new" if not res else "updated"] += 1
            if e["email"] != NA: ui.c["emails"] += 1
            if qual == "QUALIFIED": ui.c["qualified"] += 1
            if qual == "EXCLUDED": ui.c["excluded"] += 1
            ui.log(f"{badge.split()[0]} {d['name'][:38]} | {e['email']} | {round(time.time() - started, 1)}s")
        except asyncio.CancelledError:
            raise
        except Exception as ex:
            if cfg["debug"]: ui.log(f"⚠️ {type(ex).__name__}: {str(ex)[:80]}")
        finally:
            ui.c["done"] += 1
            ui.refresh()

    try:
        async with async_playwright() as p:
            context = await open_browser(p, cfg)
            search_page = context.pages[0] if context.pages else await context.new_page()
            ctx.detail_pages, ctx.render_pages = asyncio.Queue(), asyncio.Queue()
            for _ in range(cfg["tabs"]): ctx.detail_pages.put_nowait(await context.new_page())
            for _ in range(2): ctx.render_pages.put_nowait(await context.new_page())

            for loc in cfg["locations"]:
                if ctx.gate.aborted: break
                region = region_for_location(loc)
                for base in cfg["queries"]:
                    if ctx.gate.aborted: break
                    ui.log(t("log_searching", ui.lang, q=base, loc=loc))
                    urls = await collect_query(ctx, search_page, f"{base} in {loc}", cfg["max_results"])
                    if urls is None:
                        summary["blocked_queries"] += 1
                        ui.log(t("log_query_blocked", ui.lang, q=base, loc=loc)); continue
                    if not urls: summary["empty_queries"] += 1
                    new = [u for u in urls if u not in seen]
                    seen.update(new)
                    ui.c["found"] = len(seen)
                    ui.log(t("log_results", ui.lang, n=len(urls)))
                    tasks.extend(asyncio.create_task(process(u, base, loc, region)) for u in new)
                    ui.refresh()
            if ctx.gate.aborted:
                for tk in tasks:
                    if not tk.done(): tk.cancel()
            await asyncio.gather(*tasks, return_exceptions=True)
            await context.close()
    finally:
        ctx.pool.shutdown(wait=False)
        db_conn.close()
    summary["aborted"] = ctx.gate.aborted
    return run_id, summary




# ============================================================================
# 7. EXCEL EXPORT (looks like the screen)
# ============================================================================
DISPLAY_COLS = ["company", "priority_badge", "priority_score", "contactability_score", "icp_fit_score",
                "qual_status", "status", "category", "location", "email", "email_confidence", "phone",
                "whatsapp", "wa_link", "website", "maps_url", "linkedin", "facebook", "instagram",
                "matched_keywords", "email_source", "address", "enrichment_level", "enriched_at",
                "last_contacted_at", "sources_searched"]
LINK_COLS = {"website", "maps_url", "linkedin", "facebook", "instagram", "wa_link"}
BADGE_FILL = {"🔥": "FFE0CC", "🟢": "D9F2E1", "🟡": "FFF4C2", "🔴": "F8D7DA"}
QUAL_FONT = {"QUALIFIED": "1E7B45", "NOT_ICP": "7A7A7A", "EXCLUDED": "B3261E"}
CONF_FILL = {"HIGH": "D9F2E1", "MEDIUM": "FFF4C2", "LOW": "F8D7DA"}


def pretty_location(v):
    return ", ".join(p.strip().title() for p in str(v or "").split("|") if p.strip())


def display_frame(df, debug=False):
    cols = [c for c in DISPLAY_COLS if c in df.columns] + (["debug_logs"] if debug and "debug_logs" in df.columns else [])
    out = df[cols].copy()
    if "location" in out: out["location"] = out["location"].map(pretty_location)
    for c in cols:
        if c not in ("priority_score", "contactability_score", "icp_fit_score"):
            out[c] = out[c].fillna("").replace("Not available", "")
    return out


def build_excel(view, lang, kpis):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.utils import get_column_letter

    labels = COLUMN_LABELS[lang]
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospects" if lang == "en" else "Prospectos"
    teal, white = "0093A3", "FFFFFF"
    thin = Side(style="thin", color="DDE3E8")
    cols = list(view.columns)
    ws.append([labels.get(c, c) for c in cols])
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = Border(bottom=thin, right=thin)
    ws.row_dimensions[1].height = 32

    for r, row in enumerate(view.itertuples(index=False), start=2):
        zebra = PatternFill("solid", fgColor="F6F8FA") if r % 2 == 0 else None
        for i, (c, v) in enumerate(zip(cols, row), start=1):
            cell = ws.cell(row=r, column=i, value=v if v != "" else None)
            cell.border = Border(bottom=thin, right=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=c in ("matched_keywords", "address", "sources_searched"))
            if zebra: cell.fill = zebra
            if v in ("", None): continue
            if c in LINK_COLS and str(v).startswith("http"):
                cell.hyperlink = str(v)
                short = host_of(str(v)) or str(v)
                cell.value = "Google Maps" if c == "maps_url" else "WhatsApp" if c == "wa_link" else short
                cell.font = Font(color="0563C1", underline="single")
            elif c == "email":
                cell.hyperlink = f"mailto:{v}"
                cell.font = Font(color="0563C1", underline="single")
            elif c == "priority_badge":
                fill = BADGE_FILL.get(str(v)[:1])
                if fill: cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(bold=True)
            elif c == "qual_status":
                cell.font = Font(bold=True, color=QUAL_FONT.get(str(v), "333333"))
            elif c == "email_confidence" and str(v) in CONF_FILL:
                cell.fill = PatternFill("solid", fgColor=CONF_FILL[str(v)])
            elif c == "company":
                cell.font = Font(bold=True)

    last = ws.max_row
    for c, top in (("priority_score", 100), ("contactability_score", 40), ("icp_fit_score", 60)):
        if c in cols and last > 1:
            col = get_column_letter(cols.index(c) + 1)
            ws.conditional_formatting.add(f"{col}2:{col}{last}", DataBarRule(
                start_type="num", start_value=0, end_type="num", end_value=top, color="5FC4CF"))
    widths = {"company": 32, "priority_badge": 15, "priority_score": 11, "contactability_score": 15,
              "icp_fit_score": 11, "qual_status": 13, "status": 16, "category": 24, "location": 22,
              "email": 32, "email_confidence": 13, "phone": 18, "whatsapp": 16, "wa_link": 12, "website": 26,
              "maps_url": 14, "linkedin": 22, "facebook": 22, "instagram": 22, "matched_keywords": 40,
              "email_source": 16, "address": 40, "enrichment_level": 12, "enriched_at": 18,
              "last_contacted_at": 18, "sources_searched": 40, "debug_logs": 40}
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 18)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(last, 1)}"

    s2 = wb.create_sheet("Summary" if lang == "en" else "Resumen")
    s2["A1"] = "Universal Prospecting Platform"
    s2["A1"].font = Font(bold=True, size=16, color=teal)
    s2["A2"] = "Código Origami · Alejandro Moreno"
    s2["A3"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    s2["A2"].font = s2["A3"].font = Font(color="666666")
    for i, (k, v) in enumerate(kpis, start=5):
        s2.cell(row=i, column=1, value=k).font = Font(bold=True)
        s2.cell(row=i, column=2, value=v)
    s2.cell(row=5 + len(kpis) + 1, column=1, value=t("showing", lang, n=len(view)).replace("**", ""))
    s2.column_dimensions["A"].width = 34
    s2.column_dimensions["B"].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================================
# 8. USER INTERFACE
# ============================================================================
st.markdown("""
<style>
    .main { background-color: #FFFFFF; color: #31333F; }
    h1, h2, h3 { font-family: 'Segoe UI', sans-serif !important; color: #0093A3; }
    .stButton button, .stDownloadButton button { background-color: #00E5FF !important; color: #000 !important;
        font-weight: bold; width: 100%; border: none !important; }
    .stButton button p, .stDownloadButton button p { color: #000 !important; font-weight: bold; }
    .stButton button:hover, .stDownloadButton button:hover { background-color: #00B8CC !important; }
    .stButton button:disabled { background-color: #D5DBE0 !important; }
    h1 { color: #0093A3 !important; }
    div[data-testid="stMetric"], div[data-testid="metric-container"] { background-color: #F0F2F6; border-radius: 8px;
        padding: 10px; border-left: 4px solid #0093A3; box-shadow: 0 2px 4px rgba(0,0,0,0.1); color: #31333F; }
    .origami-sub { color: #5b6b78; margin-top: -8px; }
</style>
""", unsafe_allow_html=True)

LOGO_B64 = "data:image/jpeg;base64,/9j/4AAQSkZJRgABAQEASABIAAD/2wBDAAoHBwgHBgoICAgLCgoLDhgQDg0NDh0VFhEYIx8lJCIfIiEmKzcvJik0KSEiMEExNDk7Pj4+JS5ESUM8SDc9Pjv/2wBDAQoLCw4NDhwQEBw7KCIoOzs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozs7Ozv/wAARCACZAJYDASIAAhEBAxEB/8QAHAAAAQQDAQAAAAAAAAAAAAAAAAMEBQYBAgcI/8QAQhAAAgEDAgMFBQQGCQQDAAAAAQIDAAQRBSEGEjETQVFhcQcUIoGRFTJCoSMzUmKxwSRDU3KCkqLR8BY0RMJUY+H/xAAZAQACAwEAAAAAAAAAAAAAAAACAwABBAX/xAAqEQADAAICAQMDBAIDAAAAAAAAAQIDERIhMQRBURMi8CNhcbEzQqHB8f/aAAwDAQACEQMRAD8A7JRRRUIFFFJXV1BZWsl1dSrDDEpZ5HOAoHfUIK1F65xLo3Dlv2+q38duD91Tlmb0UbmuW8Ze2iaR5LHhhRHHurXsi/E39xT09T9BXKrm+u7ydp7q5lnlc5Z5HLEn1NMnHvyC6+Dr2ue3WJCY9B0xpP8A7rv4R8lB/iapeoe1bjHUcqdT90Q/htowmPn1/OqiZpT1kb61gyMy8pYkU1RKA5USdxrWqXXM1xr97M2SPjnkbI+ZpsbiYSbanKevxczd3T60jFFC5IJcHPQDOBSpsowPid0PcDGTmj0KdafbY8ttc1WzkHu/EF9DjG8c8g/nU9Y+1ri7TXC/aKX0YxtcRA5+Ywfzqp+6osfNJ2ikDccvn40mYMoHRgRjJyRkVTlP2Lm+/J2bQfbpYzYi13T3tmO3bW3xp6kHcfLNdI0jXNM160F1pd7FdRHvQ7j1B3Hzryd2TbdN/OnWnahqGj3q3en3klrOm4eN8fXxHkaW8fwNVnraiuXcF+2CG/KWPEojtbhsLHdJskhP7Q/CfPp6V1EEEAg5B6EUlpryGmn4CiiiqLCiiioQKKK0nmitoHnnkWOKNSzuxwFA6k1CDbVtVstE02bUdRnWC2hGXc/kAO8nwrzzxx7Qb3i+7MZR4NPjY9jAHI5vBn8T/Cs+0Pju44u1fkt2aPTbZiII8n4zn75HifyqomaUjBcnPnT4nXbFU2+kZLx7focY6/F1rYyQFdrcjz5zWvvExyO0benNpbXd26wQJJK5PwogyT6CmOlK2wFDf/rGR6nAwKKfzW11CxSQOrLsVYYIpmyFTQqprww3NLygQsx5e0KjOTk/nSnZEhT7yme7JO1IjbPmKeW0jNGsYZsnOwdRsAfEbUxCrbXaEjESp/pUZGMkcx/2rTsB/bxfU/7U/YsyKRI2CSP1qYzSTDmKgux3JwZF781ehatjMRDmI7Remc+NbrblhkOufAU4RnI/WHv/ABr0q1cJ6BHqF0txfPyWkbZPNuJCBnHpgZJ8PWs+fNOGHTNOGLy1xRTZraSDHaKRkZGR1FdA9nftOl4dkj0rV5ZJtLJwjsMtb+nivl3d1RXHk6y8SzhCEWBVjVTgYwo2x3b5qpzEFFwflkH+VVhr6uNVS8hZFwyOUeuopY54UmhkWSORQyOpyGB6EGt64b7JvaA2mXUfD2qTD3KYn3eVz+qc9F/un8jXcqClphJ7CiiihLCuQe2njFowvC9lJgsokvGU746qn8z8q6hrurwaDol3qlycR20ZbHiegHzJA+deVNQvp9T1C4vrpy81xIZHJ8SaZjW3sGmNwMg+VKwu6cxVlHiGAOaSBIBx39aBT0Ka2tDqF5JWxtg93KKvOm6edD0WO/fm95uSOVE2bk6hfLm2JPcAB+Kqtw61hHqUcuosRbx/Gyhcl8bhfmdqsF/xtI8znT4fd2Y/r3PNKB4KeiDyA+dc/wBS8l1wldG3BMRPJvsdaprN9ZarPHdusirJziC6IblUkNy4bONjjuqrcQ6fHZak627F7dwJIWPehGR/t8ql9cmkuorC9JVmubVecsASzISnr+AfWk7hPtLhkMQGudObDFe+Jj/6tj/PV41wmbX8MCrd3UP+UVToaWRiFiUZ3JPUCk3GDRBjtlyRjz6V0EzHS6FmWWVmVc7DJHMDWDbnBHZsGGxPMMZHWtpGHZnkMY+YyfypxpljcX93Ha26JJJKcAf86Y61LpStsCFVdId6JorX9wWcGK3hXnnlJBCL4+vcPM1dbCft0s4Eh93hnlEMC/CSkSkc7dd8kb7dxqDuRHHDHo2m8siBwskhiYmaU9/d8I6DyJPfTqJuW5vp1EawadaSLEQpxnHJzLvsC75rk+ol5FtnTwXwfFFW1u4m1DVLi7KnNxMzjfxP/wC1GNDMdim+fLO9KzxhjvMnpk0g0WDtKhyfGunjnjKRhqnVNgsEpBZV2Xqc9K9Eey/iqXiHh4W1+w+0LHCOScmRMfC/8j5jzrzv2Y/tV+tT/A/EJ4V4ptNQMmbcns7hQeqHY/TY/Kpc7RJp7PT9FYVgyhlIIIyCO+isw85X7c9c7DSbLRI2w9zJ28oH7C5AB9Sc/wCGuJVd/a/qJvuP7mINlLSNIV36bcx/NjVIrRC0hb8ikJdWJSTsyB1zilD2xIdrlSeoPabj/maQUZz+VKrGvL8QbPkaZsW17igknxyduuBt9/Y1gAq2C6n0Oa1KRjPwv5bisfAOgbPnUfYK68FrjdLrg+AMqM9rdMuWflwrqCPzRqS0W7itNQBkki93lDRzRF88yNsTv3jqPMCjh7+kaVqdtiNuWJZwJOnwHB/JjUaOVWDZs8jPc1ZYnkqhjMt8Km0I6zp0mm6hNbSDeNtj+0OoPoRg/OmEYbnBVuVh0OcYq0apGNT0OC/EivPakW82O9dyjfTK/wCEVXlhfGOXI9KPFfWq8oPJO+58M1SOSVwpJdiRjDZq0sq6BYNZRx5v7kctwysFMCn+ryehP4vp41pptoNCs11OeINeS72sRIHIP7Q5/wBPmM928XO5uXYtbczsd2M+cknOaH/K9+yBp/SWl5ZIafmC2udSkEwaKPs4A0oPNI2Rt6DmPrjxrJeO24UuniUKbmdIsknmKjLkHbx5KS1KSO2920yOFIxbqTIRMR+mOObPpgL8qNdYQ6HpVqp+Jke4cc3NuzYG/wDdUfWhtb1+7/oLG2t/sv8Allac5NakUpykmnljo95qLlLWB5GVSzYGygdST3Cn1cyu2DMun0R1FbOvI2DWtF5IelfZhrh1zgeyeRszWo92k3ySUAAJ9RiiqV7BtS31bS2b9i4Rf9Lf+tFZ6WmMXg5nxRctecV6tcsc9peSkenOcflUXTnU8/at3nr275/zGm1aELFYMc4z0q96fonDut2itazz21wF/SQsBJj94dMj+FUJGKE7kZHh1p9Y381pPHNBI6SIchl2wazeoxXa3D0x+DLMP7ltE5f8Ke7sRHdwMcZCyZibHj8W351GXOi31spZ7NuQf1irzL/mG1We11l9bhWFXjS5A3t5VDRXBz1UY+F/Lv7iOlNpLuG0u+VHk0y82VnjJKZz4bMMY8/Ss2G86+19/n58jc04PKX5+fwR/CpePWEt/wD5aPb7/vqVH5kUxYoMjnjKhjkAEFvlVj96urcpeTW1ve9k4cXEajc5yPiUAgjbZhUfrlp2Gs3RhWZ4+fMRTpykcw/I1pw0/qPfuY/USnjTXsxXhiKG5uri2kl/QXFu4l/RkdkAMhvDZgPX51JQSWdtayTPaxjTYRyJHOgDXMnjn7y92cdNh37t7exc2yaapdHmT3i9mZt4kA5gpHfsA2O8kDuqK1a/lubgJBHHDaxgpHCzBsAeO33jnJ8TSnjWbI3vocsjw4kn5Eb67mv7tpp5rRnZ+uwAA6Y8APCl9KhELtqEjWxjtRzAhP6w7J+e58lNRqRyOOYQqC+2A4A7/PxqSvVW1tYNPTso5EXtbgcpwXI+70O4U49Sa20kkoRhlvbtvx/YgD29yha6iaSZuZj2Wevjtv30txQhOsvbj/xkjgwPFECn8waX4at+11+0iExKNcA7Rgc2DnH0BqSvXtbC8mupEjuNQmlLt2v6uIk5PjzHfqdh51ly3rKkvg14EvpOqfuROmcN8yxz37NBHJjsowuZJs9MDuH7x29af69rEWlWL6NpyJAXx7yY2B7vuc34sHqfHYbDfGoaqdPie6aRn1K6G3O3MbdfHPex2x4D8qZNKXYkmlRjeWuVdpGh5FjnUrTZpI3MxNa0ZoromQtPs+1htE16e5RyvNbMmf8AEp/lRVf08uLhuTry/wAxRS6nbCTHXEtubTijVbcjHZXkq/RzUchjwecMT3YNXH2taf7h7Qb1guFuVSdfPIwfzBql0c+AWhQo5CnA3G29KKlwgIBK43wDikAcHNOQGkiVsE5O+EHiaJC62jeKS4Ug8xxsc83lmrJaalDqsYt9SfsbkAKl4AMnuAk8R+91HmOld5QSfhbPU4iX/eke1ZWwcD0ApeTEn2vJePJ/q+0WOeC90e7z2lzE6gt2isDzr3EEdVP0NPLuM6pJprtz9pcqsb/DyoCGKYwNhty0w0XWIGaCz1SL3iyD9M4aPPUqf5dDVjt7CRXtbZkLPZaiysExgqVBHqPgz9a5+bK5f3Lv5NuHCmvtfXwMdR1hrfUibSGS0iSfKwhAwYnIPMGGWJG3hg0yvdNtZIvtKxFw1vzhTArhuxbOME+B7j3+orb7a1K2Z4ve+eLuRmVl9MNn0p5pGqrHdq32Vavzgq/Yg5dWwOUqpwfTHn3U2VeKeSQmqjNblvyRWn2Kh2uZkmmS1wQrqGV2P3UPqR08AaUgtb2/X3q4vWSMsWMk6A825yAOrHvwPyqU1a4s9OBsrC2gUxyNzmWUlQ+3QYGcdN/PaoC7mub2cyzzwSs43y58AD8vKjx1kyrl4TF5Jx4vs8tE2t5Z6fpc/YxM0rkxJclscxOecqo2BwcePxdaZRTx6darqE3aO7hvdoJCCCenOR0I/iR4Zp2Y47OziN3lbezUBlVyO3mb4io27hyg+GPEiqvqN62oXTzzTqGc9ApwoHQDyxgUnFP1G0vHuzRkahJ0vHhCV3fNdTyTT5lkkbmZmY5J76a88YJPZZ8MsayyJy5EwJ8MGsBIz/XAfI10UtLSMjafb/7MF0Of0Y36bnak6UKJzYEoIx1wawUXk5u1XP7ODmrLTRYuBNJfWdbmt0XmK2zPjH7yj+dFXn2DaaGuNW1Rl+6iW6H1PM38ForPddjUuhx7dtEL22n67GP1be7S48DllP1B+tcZr1dxPocXEfDl5pMu3bp8DfsuDlT9QK8rXNtLZ3UttOhSWFyjqe4g4Io8b60VQnit+1IULyrsfChObBUNgHr506b3ho8M7kDuwMUzYqhv2zHP6NN9vu9KxzsTzco+lOSs6BZD2iq3RsYB/wB60cyNnLsSe8jrU5L5K0/gxDKQ/d18K6ZoU32jp1s4ZhK3KWKkZZojy75/dkX865gFINP7XVb+yjaO2u5oUf7yxyFQfXFYfV4Hmn7X2bfTZVir7l0TpsVtn5tUeK2TPMI+zDSn/D3Z88Cm8/ES2qNBpVv7pC2xk6yuOu7d3oMCoGSd5GZnYknqTvmksdN8+VMjC2v1HsVWRJ/prRZoOKJrkCPUYFvUXYM4xIno43+uR5U+tfslryO7S9ZbaIh3t54xzlRvyggYbOMd3WqaNsHm8O7pWS7Y2JqV6da1L0VOV73S2S+v6tc6zePMV5YwT2aDooz/AB8T41CG3mJPwHI6+VBDGtSredOxxOOeKAurutsDDLzlOQ8w6isGGQMVKnK9RWeRqwUambQOqNWRlJBUjHXyrWtiCKmuDuHn4n4ntNLXIjduaZh+GMbsf5fOqbCR3X2VaI2i8DWhkGJb3+lN6MBy/wCkCirhGiRRrHGoVEAVVHQAUVmb2xptXFPbRwg1tfLxNZxDsJ8JdBfwv0DehG3qPOu103v7G21OxmsryJZredSkiN0Iq5ensprZ5t4N0ax1O4vLrUu0ay0+1a5ljiOHkwQAoPdknrU3pmncOcR6knuun3OnwWkEtxeIJ+fnRACOXIyCehprq2nar7LuKXMSpPaTo6xGVeZJ4T1Rx49Mimn/AFxNDfWt1p2mWGni25/0cMRxIG2YOSSWBG2O6quab2gpaS7LGbTSOK9Je4tLa5tTpTxJ2UlwZEaFm5cDP3SPKnF1w7wxecRahw5bafPaXMCuYLkTlwzKucMD3elVa745mmtRa2enWWnwNMs0qWyEdsynI5iSdge7pTi99o15cPdz2+nWNpdXqlJrmJGMjKeoBJOM+VL4WHzkl7fg/R3NlczrKttFo32hdLG3xSsGYYBPTOBSFjw9pPE95py2GlXOmQTTSJNKZu0RlReYhc782AfKoS3471G2ubKaJIf6Jae5mNl5lmjySQ4PXOaUn4+vALNNOtLXTY7Oczxpbht3IwSSxOQRtipwsrlJJHSuG9f07VBpOn3Fhc6dA1wjyTlxLGv3gwPQ48KZ8P6NpMXDl3xBq1vJdpFOtvDbJJyBnIySxG+APCm9/wAbzXlhdWlppljp4vMG5ktoyGl3zjJJwM9wpponFNxotvcWb20F5ZXWO1t7hSVJHQgjBBHlR8b0DyjZcbDhLQdQ1DTL62s52sr+2uHNm0hLLJGucBhuQcjFKwcG6PdX+htNpV1pvvtxJFLZTyNzMqrkOM4Yb7VWG9oOofaMF1HBbwx21u9vBbxKVSJWBBI3znfr5Unbce6hBNpk8kcU8+mFuxlkyWZTn4WOdwMnFBwsPlBMHRtA0PS7K51HT5tRn1KSTs40mMaxRq3L3dWJ+VZ0PhXS5Nb1SfVrO6g02xYL2Dv+kDO2FUkd4GSfSoax45nt7NLW70+z1CGGUywC5UkwsTk8pBG2e6trj2ia1Kkwt5vcpLi4a4mltmZGkYjGDv0A7qnCycoJ7R+A7OXV9f0m/YrLZoBbSlsAMzYQnyOR9ab/APRFtY8HXN9qSMupdvGEhJIMcZfkJI8yG+lRFxx/qNzDcLJHE01zapbSz/FzuFOQxOfvbdab3PG+qXkV6t46zvemEySMMEdn90DG3rV8LK5QI8daVaaNxbf2FkhS3hcBFLEkDlB6n1rr3sh4QbQdCbVL2LlvtQAYA9Y4vwjyJ6n5eFVbgvhu49oHFE3Fes2yJYiTm7NchZpAAAAD+EdT9PGu1gYGBTW2loXrvYUUUUBYUUUVCEXxFw9p/E+kS6bqMZaN91dfvRt3Mp7jXnHi7hDUuD9TNteJzQSEm3uF+7Ko/gfEV6ipnquk2Ot6fJYajbpPBIMFWHQ+I8D50c1oprZ5K5jRmui8Z+yHU9FeS80RX1CxGWKAZliHmPxDzH0rnJBBIIwR3Gnpp+BbWgzRmigDOfIZqyheDOd15sjb4sd//Nqw8c2x7N9xtkdc+FZiikYK1ujHPwnJB37qBDdPsA5wPH5VfsL2t72jQwzAZMbfSjsJtv0bbjI2rYx3PIc83KCQfi2z30NFdRAuedQvfn/njU0Xy/dGrwzR/eQjfFasjoAWUgHpkUGWQ7mRj863hjuLuZLeBJJpJGASNAWLHuAHearoJcvcSzV39n3s6uuLblby8DwaTG3xyjZpSPwr/M91WXgn2NSyMmocUDs0BDJZKcs3989w8h+VdjggitoUggiSKKMBURFAVR4ADpSqv2QxT8mlpaW9haRWlpCkMEShUjQYCgUtRRSQwoooqECiiioQKKKKhAqrcS+zrh3icNJc2vu10d/ebYBHJ89sH51aaKtPRDg2u+xLXbFi+kTxalF15TiKQfInB+tUy94a1vSGb7R065swBjnliIT/ADdK9WVrL+qb0NGsj9wHOzyEse+O0UD1rIjbGDKo2zgtXQ/aB/3EnrXOf63505VsByxYwLy5FyrDbYA5yetSencH8Sauyix0i6nRuknIQh/xHAroPs0/7mL1FdnoKya9i5j5ZxDQPYdqNzyy67fJZJ3wwASOR69B+ddS4d4M0LheELptkolxhriQBpW9W/2xU7RSnTYxJIKKKKEsKKKKhAoooqEP/9k="

if "lang" not in st.session_state: st.session_state.lang = "en"
col_logo, col_title, col_lang = st.columns([1, 7, 2])
with col_logo:
    st.image(LOGO_B64, width=80)
with col_lang:
    choice = st.radio("🌐", ["English", "Español"], horizontal=True,
                      index=0 if st.session_state.lang == "en" else 1, key="lang_radio",
                      help="Interface language / Idioma de la interfaz")
    st.session_state.lang = "en" if choice == "English" else "es"
L = st.session_state.lang
with col_title:
    st.title("Universal Prospecting Platform")
    st.markdown(f"<div class='origami-sub'>{t('subtitle', L, v=__version__)}</div>", unsafe_allow_html=True)

tab_discover, tab_qualify, tab_prospect = st.tabs([t("tab1", L), t("tab2", L), t("tab3", L)])

# ---------------------------------------------------------
# TAB 1: DISCOVER
# ---------------------------------------------------------
with tab_discover:
    st.markdown(f"### {t('icp_title', L)}")
    col_icp1, col_icp2 = st.columns(2)
    with col_icp1:
        st.markdown(f"#### {t('disc_params', L)}")
        target_locations = st.text_area(t("locations", L), "Santiago, Chile\nBogota, Colombia", height=90,
                                        help=t("locations_help", L))
        base_keywords = st.text_area(t("queries", L), "used car dealer\nauto parts importer", height=90,
                                     help=t("queries_help", L))
        max_results = st.selectbox(t("max_results", L), [20, 50, 100, 120, 200], index=1, help=t("max_results_help", L))
        modes = {"FAST": t("mode_fast", L), "SMART": t("mode_smart", L), "DEEP": t("mode_deep", L)}
        enrichment_mode = st.radio(t("mode", L), list(modes), index=1, format_func=lambda k: modes[k],
                                   help=t("mode_help", L))
        with st.expander(t("speed_box", L)):
            s1, s2 = st.columns(2)
            with s1:
                tabs_n = st.slider(t("tabs", L), 1, 4, 2, help=t("tabs_help", L))
                refresh_days = st.number_input(t("refresh_days", L), 1, 365, 30, help=t("refresh_days_help", L))
                debug_mode = st.checkbox(t("debug", L), help=t("debug_help", L))
            with s2:
                fast_browser = st.checkbox(t("fast_browser", L), value=True, help=t("fast_browser_help", L))
                show_browser = st.checkbox(t("show_browser", L), value=False, help=t("show_browser_help", L))
                verify_mx = st.checkbox(t("verify_mx", L), value=HAS_DNS, disabled=not HAS_DNS, help=t("verify_mx_help", L))
                search_engine = st.checkbox(t("search_engine", L), value=True, help=t("search_engine_help", L))
                social_search = st.checkbox(t("social_search", L), value=True, help=t("social_search_help", L))
    with col_icp2:
        st.markdown(f"#### {t('icp_box', L)}")
        req_kws_input = st.text_area(t("req", L), "importer | wholesale\nvehicles | cars | auto", height=80, help=t("req_help", L))
        st.caption(t("req_caption", L))
        strong_kws_input = st.text_area(t("strong", L), "japanese | japan: 25\nused | second hand: 20", height=70, help=t("strong_help", L))
        bonus_kws_input = st.text_area(t("bonus", L), "b2b:10\ndealer:5", height=60, help=t("bonus_help", L))
        exclude_keywords = st.text_area(t("excl", L), "rental, repair | mechanic, motorcycle", height=60, help=t("excl_help", L))

    locations = [x.strip() for x in target_locations.splitlines() if x.strip()]
    bases = [x.strip() for x in base_keywords.splitlines() if x.strip()]
    per_lead = {"FAST": 2.5, "SMART": 4.5, "DEEP": 9}[enrichment_mode] / max(1, tabs_n * 0.8)
    searches = len(locations) * len(bases)
    est = int((searches * (12 + min(max_results, 120) * per_lead)) / 60) + 1
    st.caption(t("estimate", L, m=est, s=searches))

    if st.button(t("start", L), help=t("start_help", L)):
        if not locations or not bases:
            st.error(t("need_input", L)); st.stop()
        cfg = {
            "locations": locations, "queries": bases, "max_results": int(max_results), "mode": enrichment_mode,
            "tabs": int(tabs_n), "fast_browser": fast_browser, "headless": not show_browser,
            "verify_mx": verify_mx and HAS_DNS, "search_engine": search_engine, "social_search": social_search,
            "refresh_days": int(refresh_days), "debug": debug_mode,
            "req": parse_flat_keywords(req_kws_input), "strong": parse_weighted_keywords(strong_kws_input),
            "bonus": parse_weighted_keywords(bonus_kws_input), "excl": parse_flat_keywords(exclude_keywords),
        }
        _ddg_state["blocked"] = False
        started = time.time()
        with st.status(t("running", L), expanded=True) as status_box:
            ui = RunUI(L)
            try:
                run_id, summary = asyncio.run(run_discovery(cfg, ui))
            except Exception as e:
                msg = str(e)
                if "ProcessSingleton" in msg or "user data directory is already in use" in msg.lower():
                    st.error("Chrome profile in use: close other runs of this app and try again." if L == "en"
                             else "El perfil de Chrome está en uso: cierra otras ejecuciones de la app y vuelve a intentarlo.")
                else:
                    st.error(f"{type(e).__name__}: {msg[:300]}")
                st.stop()
            status_box.update(label=t("done", L, m=round((time.time() - started) / 60, 1)), state="complete")
        if summary.get("aborted"): st.warning(t("aborted", L))
        else: st.success(t("done", L, m=round((time.time() - started) / 60, 1)))
        if summary["no_identity"]: st.warning(t("warn_no_identity", L, n=summary["no_identity"]))
        if summary["conflicts"]: st.warning(t("warn_conflicts", L, n=summary["conflicts"]))
        if summary["blocked_queries"]: st.warning(t("warn_blocked_q", L, n=summary["blocked_queries"]))

# ---------------------------------------------------------
# TAB 2: QUALIFY
# ---------------------------------------------------------
with tab_qualify:
    st.markdown(f"### {t('dash_title', L)}")
    db_df = load_leads()
    if db_df.empty:
        st.info(t("empty_db", L))
    else:
        for col in DISPLAY_COLS + ["debug_logs"]:
            if col not in db_df.columns: db_df[col] = None
        is_na = lambda s: s.isna() | (s == "Not available") | (s == "")
        kpi = {
            "m_total": len(db_df),
            "m_qual": int((db_df["qual_status"] == "QUALIFIED").sum()),
            "m_noticp": int((db_df["qual_status"] == "NOT_ICP").sum()),
            "m_excl": int((db_df["qual_status"] == "EXCLUDED").sum()),
            "m_email": int(((db_df["qual_status"] == "QUALIFIED") & ~is_na(db_df["email"])).sum()),
            "m_wa": int((~is_na(db_df["whatsapp"]) | ~is_na(db_df["wa_link"])).sum()),
            "m_contacted": int((db_df["status"] == "Contacted").sum()),
            "m_replied": int((db_df["status"] == "Replied").sum()),
        }
        keys = list(kpi)
        for row_keys in (keys[:4], keys[4:]):
            cols_m = st.columns(4)
            for cm, k in zip(cols_m, row_keys):
                cm.metric(t(k, L), kpi[k], help=t(k + "_help", L))
        st.markdown("---")
        st.markdown(f"#### {t('filters', L)}")
        f1, f2, f3, f4 = st.columns(4)
        with f1:
            icp_filter = st.multiselect(t("f_icp", L), ["QUALIFIED", "NOT_ICP", "EXCLUDED"], default=["QUALIFIED"], help=t("f_icp_help", L), placeholder=t("choose", L))
            loc_options = sorted({pretty_location(x) for x in db_df["location"].dropna() if x})
            loc_filter = st.multiselect(t("f_loc", L), loc_options, help=t("f_loc_help", L), placeholder=t("choose", L))
        with f2:
            quality_filter = st.multiselect(t("f_badge", L), ["🔥 EXCELLENT", "🟢 GOOD", "🟡 NORMAL", "🔴 LOW"],
                                            default=["🔥 EXCELLENT", "🟢 GOOD", "🟡 NORMAL"], help=t("f_badge_help", L), placeholder=t("choose", L))
            search_txt = st.text_input(t("f_search", L), placeholder=t("f_search_ph", L), help=t("f_search_help", L))
        with f3:
            req_email = st.checkbox(t("f_email", L), help=t("f_email_help", L))
            req_phone = st.checkbox(t("f_phone", L), help=t("f_phone_help", L))
        with f4:
            min_score = st.number_input(t("f_min", L), min_value=0, max_value=100, value=0, help=t("f_min_help", L))

        fdf = db_df.copy()
        if icp_filter: fdf = fdf[fdf["qual_status"].isin(icp_filter)]
        if quality_filter: fdf = fdf[fdf["priority_badge"].isin(quality_filter)]
        if loc_filter: fdf = fdf[fdf["location"].map(pretty_location).isin(loc_filter)]
        if req_email: fdf = fdf[~is_na(fdf["email"])]
        if req_phone: fdf = fdf[~is_na(fdf["phone"]) | ~is_na(fdf["whatsapp"]) | ~is_na(fdf["wa_link"])]
        fdf = fdf[fdf["priority_score"].fillna(0) >= min_score]
        if search_txt.strip():
            q = search_txt.strip().lower()
            hay = (fdf["company"].fillna("") + " " + fdf["category"].fillna("") + " " + fdf["email"].fillna("") + " " +
                   fdf["website"].fillna("") + " " + fdf["matched_keywords"].fillna("")).str.lower()
            fdf = fdf[hay.str.contains(re.escape(q), na=False)]

        view = display_frame(fdf, debug=debug_mode) if not fdf.empty else pd.DataFrame(columns=DISPLAY_COLS)
        lab = COLUMN_LABELS[L]
        cc = st.column_config
        config = {c: cc.TextColumn(lab.get(c, c)) for c in view.columns}
        config.update({
            "priority_score": cc.ProgressColumn(lab["priority_score"], min_value=0, max_value=100, format="%d"),
            "contactability_score": cc.ProgressColumn(lab["contactability_score"], min_value=0, max_value=40, format="%d"),
            "icp_fit_score": cc.ProgressColumn(lab["icp_fit_score"], min_value=0, max_value=60, format="%d"),
            "website": cc.LinkColumn(lab["website"]),
            "maps_url": cc.LinkColumn(lab["maps_url"], display_text="📍 Maps"),
            "wa_link": cc.LinkColumn(lab["wa_link"], display_text="💬 WhatsApp"),
            "linkedin": cc.LinkColumn(lab["linkedin"]), "facebook": cc.LinkColumn(lab["facebook"]),
            "instagram": cc.LinkColumn(lab["instagram"]),
        })
        try:
            config["company"] = cc.TextColumn(lab["company"], pinned=True)
        except TypeError:
            config["company"] = cc.TextColumn(lab["company"])
        config = {k: v for k, v in config.items() if k in view.columns}
        st.dataframe(view, use_container_width=True, hide_index=True, column_config=config, height=520)
        st.markdown(t("showing", L, n=len(fdf)))

        b1, b2, b3 = st.columns(3)
        with b1:
            if not view.empty:
                kpis = [(t(k, L), v) for k, v in kpi.items()]
                st.download_button(t("excel", L), data=build_excel(view, L, kpis),
                                   file_name=f"origami_prospects_{datetime.date.today().isoformat()}.xlsx",
                                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                   help=t("excel_help", L))
        with b2:
            if not view.empty:
                csv_view = view.rename(columns=lab)
                st.download_button(t("csv", L), data=csv_view.to_csv(index=False, sep=";").encode("utf-8-sig"),
                                   file_name="origami_prospects.csv", mime="text/csv", help=t("csv_help", L))
        with b3:
            if st.button(t("recover", L), help=t("recover_help", L)):
                l_res, a_res = reset_stuck_sends()
                st.success(t("recovered", L, a=l_res, b=a_res))
                time.sleep(2); st.rerun()
        with st.expander(t("danger", L)):
            sure = st.checkbox(t("confirm_clear", L), help=t("confirm_clear_help", L))
            if st.button(t("clear", L), disabled=not sure, help=t("clear_help", L)):
                if clear_db(): st.rerun()
                else: st.error(t("clear_blocked", L))

# ---------------------------------------------------------
# TAB 3: PROSPECT
# ---------------------------------------------------------
with tab_prospect:
    st.markdown(f"### {t('out_title', L)}")
    st.markdown(t("out_sub", L))
    templates = sorted(f for f in os.listdir(BASE_DIR) if re.match(r"template_.*\.html?$", f, re.I))
    col_em1, col_em2 = st.columns(2)
    with col_em1:
        sender_email = st.text_input(t("sender", L), placeholder="youremail@gmail.com", help=t("sender_help", L))
        sender_password = st.text_input(t("password", L), type="password", placeholder="xxxx xxxx xxxx xxxx", help=t("password_help", L))
        email_subject = st.text_input(t("subject", L), placeholder="Partnership Inquiry", help=t("subject_help", L))
        template_file = st.selectbox(t("template", L), templates, help=t("template_help", L)) if templates else None
        if not templates: st.warning(t("no_templates", L))
    with col_em2:
        target_quality = st.multiselect(t("t_quality", L), ["🔥 EXCELLENT", "🟢 GOOD", "🟡 NORMAL"],
                                        default=["🔥 EXCELLENT", "🟢 GOOD"], help=t("t_quality_help", L), placeholder=t("choose", L))
        max_emails_this_run = st.number_input(t("max_emails", L), min_value=1, max_value=200, value=20, help=t("max_emails_help", L))
        delay_range = st.slider(t("delay", L), 5, 120, (20, 45), help=t("delay_help", L))

    html_template_content = ""
    if template_file:
        with open(os.path.join(BASE_DIR, template_file), "r", encoding="utf-8") as fh:
            html_template_content = fh.read()
        with st.expander(t("preview", L)):
            st.markdown(f"**{fill_template(email_subject or '—', t('preview_company', L))}**")
            components.html(fill_template(html_template_content, t("preview_company", L)), height=420, scrolling=True)

    with st.expander(t("guide", L)):
        st.markdown(t("guide_text", L))

    c_test, c_start = st.columns(2)
    with c_test:
        if st.button(t("test", L), help=t("test_help", L)):
            if not sender_email or not sender_password or not email_subject or not html_template_content:
                st.error(t("need_creds", L))
            else:
                try:
                    srv = connect_smtp(sender_email, sender_password.replace(" ", ""))
                    ok, res, srv = send_single_email(srv, sender_email, sender_password, sender_email,
                                                     fill_template(email_subject, t("preview_company", L)),
                                                     t("preview_company", L), html_template_content)
                    srv.quit()
                    st.success(t("test_ok", L, e=sender_email)) if ok else st.error(res)
                except Exception as e:
                    st.error(t("smtp_fail", L, e=type(e).__name__))
    with c_start:
        start_campaign = st.button(t("start_camp", L), help=t("start_camp_help", L))

    if start_campaign:
        if not sender_email or not sender_password or not email_subject or not html_template_content:
            st.error(t("need_creds", L)); st.stop()
        try:
            active_server = connect_smtp(sender_email, sender_password.replace(" ", ""))
        except Exception as e:
            st.error(t("smtp_fail", L, e=type(e).__name__)); st.stop()
        st.markdown(t("camp_log", L))
        log_box = st.empty()
        campaign_run_id = f"CAMP-{uuid.uuid4().hex[:8].upper()}"
        sent_count, stop_campaign, no_more_leads, skipped_ids = 0, False, False, set()
        while sent_count < max_emails_this_run:
            campaign_targets = get_campaign_candidates(target_quality, (max_emails_this_run - sent_count) * 2, list(skipped_ids))
            if campaign_targets.empty:
                no_more_leads = True; break
            locked_in_batch = 0
            for _, lead in campaign_targets.iterrows():
                if sent_count >= max_emails_this_run: break
                email_lower = lead["email"].strip().lower() if lead["email"] else ""
                if not email_lower or email_lower == "not available": continue
                subj = fill_template(email_subject, lead["company"])
                fresh_lead = lock_lead_and_create_activity(lead["lead_id"], campaign_run_id, email_lower, sender_email, subj)
                if not fresh_lead: continue
                locked_in_batch += 1
                if fresh_lead["qual_status"] != "QUALIFIED" or fresh_lead["priority_badge"] not in target_quality:
                    release_send_lock(lead["lead_id"], fresh_lead["lock_token"], email_lower, fresh_lead.get("activity_id"))
                    skipped_ids.add(lead["lead_id"]); continue
                success, new_status, active_server = send_single_email(
                    active_server, sender_email, sender_password, fresh_lead["email"], subj,
                    fresh_lead["company"], html_template_content)
                err_msg = "" if success else new_status
                try:
                    update_activity_and_lead(fresh_lead["activity_id"], lead["lead_id"], new_status, new_status,
                                             fresh_lead["lock_token"], err_msg, email_lower)
                except RuntimeError as re_err:
                    log_box.error(str(re_err)); skipped_ids.add(lead["lead_id"])
                    if new_status in ["Send outcome unknown", "Sender rejected"]:
                        stop_campaign = True; break
                    continue
                except Exception as gen_err:
                    log_box.error(f"DB: {type(gen_err).__name__}"); stop_campaign = True; break
                if success:
                    sent_count += 1
                    sleep_time = random.uniform(delay_range[0], delay_range[1])
                    log_box.success(t("sent_line", L, i=sent_count, n=max_emails_this_run, c=fresh_lead["company"],
                                      e=fresh_lead["email"], s=round(sleep_time, 1)))
                    time.sleep(sleep_time)
                else:
                    skipped_ids.add(lead["lead_id"])
                    log_box.error(t("fail_line", L, c=fresh_lead["company"], r=new_status))
                    if new_status in ["Send outcome unknown", "Sender rejected"]:
                        stop_campaign = True; break
            if stop_campaign: break
            if locked_in_batch == 0:
                log_box.info(t("no_lock", L)); no_more_leads = True; break
        try: active_server.quit()
        except Exception: pass
        if stop_campaign: st.warning(t("camp_stop", L, n=sent_count))
        elif no_more_leads and sent_count < max_emails_this_run: st.info(t("camp_early", L, n=sent_count))
        else: st.success(t("camp_done", L, n=sent_count))
